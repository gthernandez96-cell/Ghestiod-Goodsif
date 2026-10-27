"""API v1 — Cuentas por Cobrar (AR) & Control de Créditos a Clientes 360°."""
from fastapi import APIRouter, Depends, HTTPException, Query, Response
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, and_, func, desc
from sqlalchemy.orm import selectinload
from pydantic import BaseModel
from typing import Optional, List
from decimal import Decimal
from datetime import datetime, timezone, timedelta
import uuid
import io
import csv
import os

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as OpenpyxlImage

try:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib import colors as rl_colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table as RLTable, TableStyle as RLTableStyle, Image as RLImage
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    HAS_REPORTLAB = True
except ImportError:
    HAS_REPORTLAB = False

from app.core.database import get_db
from app.core.security import get_current_user
from app.models.ar import CustomerInvoiceCredit, CustomerCreditPayment
from app.models.partner import Partner, PartnerType, CreditStatus
from app.models.company import Company
from app.models.sale import SaleOrder, SalePayment, SaleState, SaleType
from app.api.v1.partners import check_credit_payment_permission

router = APIRouter(prefix="/ar")


async def _sync_sales_to_credits(db: AsyncSession, company_id: uuid.UUID):
    """Sincronizar ventas al crédito, pre-ventas y clientes con saldo de crédito pendiente hacia Cuentas por Cobrar."""
    try:
        # 1. Obtener todas las cuentas por cobrar existentes
        credits_stmt = select(CustomerInvoiceCredit).where(
            CustomerInvoiceCredit.company_id == company_id,
            CustomerInvoiceCredit.is_deleted == False
        )
        credits_res = await db.execute(credits_stmt)
        existing_credits = credits_res.scalars().all()

        existing_numbers = {c.number for c in existing_credits if c.number}
        existing_sale_ids = {c.sale_id for c in existing_credits if c.sale_id}

        synced_any = False

        # Si alguna orden vinculada fue cancelada o rechazada, anular su cuenta por cobrar
        cancelled_sales_stmt = select(CustomerInvoiceCredit).join(
            SaleOrder, CustomerInvoiceCredit.sale_id == SaleOrder.id
        ).where(
            CustomerInvoiceCredit.company_id == company_id,
            CustomerInvoiceCredit.is_deleted == False,
            SaleOrder.state.in_([SaleState.CANCELLED, SaleState.REJECTED])
        )
        cancelled_invs = (await db.execute(cancelled_sales_stmt)).scalars().all()
        for cinv in cancelled_invs:
            cinv.is_deleted = True
            cinv.state = "cancelled"
            synced_any = True

        # 2. Sincronizar Ventas al Crédito y Pre-ventas con cliente identificado
        sales_stmt = (
            select(SaleOrder)
            .options(selectinload(SaleOrder.payments))
            .where(
                SaleOrder.company_id == company_id,
                SaleOrder.customer_id.isnot(None),
                SaleOrder.is_deleted == False
            )
        )
        sales_res = await db.execute(sales_stmt)
        sales = sales_res.scalars().all()

        for s in sales:
            if s.id in existing_sale_ids or s.number in existing_numbers:
                continue

            # Verificar si la venta tuvo pago explícito al CRÉDITO
            credit_paid = sum(float(p.amount or 0) for p in (s.payments or []) if p.payment_method and "credit" in str(p.payment_method).lower())
            
            # O si es una Pre-venta con cliente identificado (no cancelada ni rechazada)
            is_presale = (
                (s.sale_type == SaleType.PRESALE or (s.number and s.number.startswith("PRE-")))
                and s.state not in [SaleState.CANCELLED, SaleState.REJECTED]
                and (s.total or 0) > 0.01
            )

            if credit_paid > 0.01 or is_presale:
                tot_amount = Decimal(str(credit_paid if credit_paid > 0.01 else (s.total or 0)))
                bal_amount = Decimal(str(credit_paid if credit_paid > 0.01 else (s.amount_due if (s.amount_due is not None and s.amount_due > 0) else s.total or 0)))

                order_dt = s.sale_date or s.created_at or datetime.now(timezone.utc)
                credit_days = 30
                due_date = s.due_date or (order_dt + timedelta(days=credit_days))

                doc_notes = f"Pre-venta {s.number} - Cliente: {s.customer_name or 'Cliente'}" if is_presale else f"Venta al Crédito {s.number} - Cliente: {s.customer_name or 'Cliente'}"

                credit_inv = CustomerInvoiceCredit(
                    company_id=company_id,
                    customer_id=s.customer_id,
                    sale_id=s.id,
                    number=s.number or f"CRED-{str(s.id)[:8].upper()}",
                    issue_date=order_dt,
                    due_date=due_date,
                    total=tot_amount,
                    balance=bal_amount,
                    state="open" if bal_amount > 0 else "paid",
                    notes=doc_notes
                )
                db.add(credit_inv)
                existing_numbers.add(credit_inv.number)
                existing_sale_ids.add(s.id)
                existing_credits.append(credit_inv)
                synced_any = True

        # 3. Sincronizar Clientes con saldo pendiente de crédito en su cuenta (Partner.balance > 0)
        partners_stmt = select(Partner).where(
            Partner.company_id == company_id,
            Partner.balance > 0.01,
            Partner.is_deleted == False
        )
        partners_res = await db.execute(partners_stmt)
        debtor_partners = partners_res.scalars().all()

        for p in debtor_partners:
            p_balance = Decimal(str(p.balance or 0))
            if p_balance <= Decimal("0.01"):
                continue

            # Sumar cuentas de crédito abiertas de este cliente
            existing_p_bal = sum(
                c.balance for c in existing_credits 
                if c.customer_id == p.id and c.balance and c.balance > 0
            )

            diff_balance = p_balance - existing_p_bal
            if diff_balance > Decimal("0.01"):
                now_dt = datetime.now(timezone.utc)
                due_dt = now_dt + timedelta(days=p.payment_days or 30)
                inv_num = f"CRED-{str(p.id)[:6].upper()}"
                if inv_num in existing_numbers:
                    inv_num = f"CRED-{str(p.id)[:4].upper()}-{str(uuid.uuid4())[:4].upper()}"

                credit_inv = CustomerInvoiceCredit(
                    company_id=company_id,
                    customer_id=p.id,
                    number=inv_num,
                    issue_date=now_dt,
                    due_date=due_dt,
                    total=diff_balance,
                    balance=diff_balance,
                    state="open",
                    notes=f"Cuenta de Crédito - {p.name}"
                )
                db.add(credit_inv)
                existing_numbers.add(credit_inv.number)
                existing_credits.append(credit_inv)
                synced_any = True

        if synced_any:
            await db.commit()
    except Exception as e:
        import logging
        logging.getLogger(__name__).warning(f"Error sincronizando ventas a créditos: {e}")


# ===== SCHEMAS =====
class CustomerInvoiceCreditCreateIn(BaseModel):
    customer_id: Optional[str] = None
    customer_name: Optional[str] = None
    customer_nit: Optional[str] = "CF"
    number: str
    issue_date: Optional[str] = None
    credit_days: int = 30
    total: float
    notes: Optional[str] = None


class ARPaymentIn(BaseModel):
    amount: float
    payment_method: str = "cash"  # cash, card, transfer, check
    reference: Optional[str] = None
    notes: Optional[str] = None


# ===== 1. KPIS =====
@router.get("/kpis")
async def get_ar_kpis(
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Métricas ejecutivas de Cuentas por Cobrar y Cartera de Clientes unificadas por cliente."""
    c_uuid = uuid.UUID(str(current_user.company_id))
    await _sync_sales_to_credits(db, c_uuid)

    now = datetime.now(timezone.utc)
    month_start = datetime(now.year, now.month, 1, tzinfo=timezone.utc)

    # 1. Total por cobrar pendiente por clientes
    p_stmt = select(Partner).where(
        Partner.company_id == c_uuid,
        Partner.balance > 0,
        Partner.is_deleted == False
    )
    p_res = await db.execute(p_stmt)
    debtor_partners = p_res.scalars().all()

    total_receivable = sum(float(p.balance or 0.0) for p in debtor_partners)
    total_debtors = len(debtor_partners)

    # 2. Calcular vencimientos por cliente (a partir de su primera venta al crédito)
    c_stmt = select(CustomerInvoiceCredit).where(
        CustomerInvoiceCredit.company_id == c_uuid,
        CustomerInvoiceCredit.balance > 0,
        CustomerInvoiceCredit.is_deleted == False
    ).order_by(CustomerInvoiceCredit.issue_date.asc())
    c_res = await db.execute(c_stmt)
    open_credits = c_res.scalars().all()

    credits_by_cust = {}
    for c in open_credits:
        cid = str(c.customer_id) if c.customer_id else None
        if cid and cid not in credits_by_cust:
            credits_by_cust[cid] = c.issue_date

    due_soon_count = 0
    overdue_count = 0
    overdue_amount = 0.0

    for p in debtor_partners:
        p_bal = float(p.balance or 0.0)
        first_dt = credits_by_cust.get(str(p.id), p.created_at or now)
        if first_dt and first_dt.tzinfo is None:
            first_dt = first_dt.replace(tzinfo=timezone.utc)
        
        due_dt = (first_dt + timedelta(days=p.payment_days or 30)) if first_dt else (now + timedelta(days=30))
        days_diff = (due_dt.date() - now.date()).days
        
        if days_diff < 0:
            overdue_count += 1
            overdue_amount += p_bal
        elif days_diff <= 7:
            due_soon_count += 1

    # 3. Cobros y abonos recibidos este mes
    collected_stmt = (
        select(func.coalesce(func.sum(CustomerCreditPayment.amount), Decimal("0")))
        .where(
            CustomerCreditPayment.company_id == c_uuid,
            CustomerCreditPayment.payment_date >= month_start
        )
    )
    collected_res = await db.execute(collected_stmt)
    collected_this_month = float(collected_res.scalar() or 0.0)

    return {
        "total_receivable": total_receivable,
        "total_open_invoices": total_debtors,
        "due_soon_count": due_soon_count,
        "overdue_count": overdue_count,
        "overdue_amount": overdue_amount,
        "collected_this_month": collected_this_month,
    }


# ===== 2. LIST INVOICES / CREDITS (UNIFICADO POR CLIENTE) =====
@router.get("/invoices")
async def list_customer_invoices(
    status: Optional[str] = Query(None),  # 'all', 'open', 'due_soon', 'overdue', 'paid'
    search: Optional[str] = Query(None),
    customer_id: Optional[str] = Query(None),
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Listar cuentas por cobrar UNIFICADAS POR CLIENTE.
    La primera venta al crédito marca el inicio del plazo de crédito del cliente.
    """
    c_uuid = uuid.UUID(str(current_user.company_id))
    await _sync_sales_to_credits(db, c_uuid)

    now = datetime.now(timezone.utc)

    # 1. Obtener todos los clientes
    p_stmt = select(Partner).where(
        Partner.company_id == c_uuid,
        Partner.is_deleted == False
    ).order_by(Partner.name.asc())
    if customer_id:
        try:
            p_stmt = p_stmt.where(Partner.id == uuid.UUID(customer_id))
        except ValueError:
            pass

    p_res = await db.execute(p_stmt)
    partners = p_res.scalars().all()

    # 2. Obtener todas las cuentas de crédito vinculadas
    c_stmt = (
        select(CustomerInvoiceCredit)
        .options(selectinload(CustomerInvoiceCredit.payments))
        .where(
            CustomerInvoiceCredit.company_id == c_uuid,
            CustomerInvoiceCredit.is_deleted == False
        )
        .order_by(CustomerInvoiceCredit.issue_date.asc())
    )
    c_res = await db.execute(c_stmt)
    all_credits = c_res.scalars().all()

    credits_by_customer = {}
    for c in all_credits:
        cid = str(c.customer_id) if c.customer_id else None
        if not cid:
            continue
        if cid not in credits_by_customer:
            credits_by_customer[cid] = []
        credits_by_customer[cid].append(c)

    items = []
    for p in partners:
        p_id_str = str(p.id)
        p_credits = credits_by_customer.get(p_id_str, [])
        open_credits = [c for c in p_credits if (c.balance or 0) > 0]

        p_balance = float(p.balance or 0.0)
        credits_balance = sum(float(c.balance or 0.0) for c in open_credits)
        bal = max(p_balance, credits_balance)

        # Si el cliente no tiene deuda y no tiene historial, omitir
        if bal <= 0 and len(p_credits) == 0:
            continue

        tot = sum(float(c.total or 0.0) for c in p_credits) if p_credits else bal
        if tot < bal:
            tot = bal

        # Regla de negocio: la primera venta al crédito marca el inicio del conteo de días
        first_credit_date = None
        if open_credits:
            first_credit_date = open_credits[0].issue_date
        elif p_credits:
            first_credit_date = p_credits[0].issue_date
        else:
            first_credit_date = p.created_at or now

        if first_credit_date and first_credit_date.tzinfo is None:
            first_credit_date = first_credit_date.replace(tzinfo=timezone.utc)

        p_days = p.payment_days or 30
        due_date = (first_credit_date + timedelta(days=p_days)) if first_credit_date else (now + timedelta(days=p_days))

        days_diff = (due_date.date() - now.date()).days
        
        d_status = "paid"
        if bal > 0:
            if days_diff < 0:
                d_status = "overdue"
            elif days_diff <= 7:
                d_status = "due_soon"
            else:
                d_status = "on_time"

        # Filtrar por status
        if status and status != "all":
            if status == "open" and bal <= 0:
                continue
            elif status == "due_soon" and d_status != "due_soon":
                continue
            elif status == "overdue" and d_status != "overdue":
                continue
            elif status == "paid" and d_status != "paid":
                continue

        c_name = p.name
        c_nit = p.nit or "CF"
        c_code = f"CLI-{str(p.id)[:6].upper()}"

        # Número de referencia representativo
        rep_number = open_credits[0].number if open_credits else (p_credits[0].number if p_credits else c_code)

        if search:
            q_low = search.lower()
            matches_inv_num = any(q_low in (c.number or "").lower() for c in p_credits)
            if q_low not in c_name.lower() and q_low not in c_nit.lower() and q_low not in c_code.lower() and not matches_inv_num:
                continue

        rep_id = str(open_credits[0].id) if open_credits else p_id_str

        items.append({
            "id": rep_id,
            "customer_id": p_id_str,
            "number": rep_number,
            "customer_name": c_name,
            "customer_nit": c_nit,
            "customer_phone": p.phone or p.mobile or "",
            "first_credit_date": first_credit_date.strftime("%d/%m/%Y") if first_credit_date else "",
            "issue_date": first_credit_date.strftime("%d/%m/%Y") if first_credit_date else "",
            "due_date": due_date.strftime("%d/%m/%Y") if due_date else "Sin fecha",
            "payment_days": p_days,
            "days_diff": days_diff,
            "total": tot,
            "balance": bal,
            "credit_limit": float(p.credit_limit or 0.0),
            "available_credit": max(0.0, float(p.credit_limit or 0.0) - bal),
            "status": d_status,
            "invoices_count": len(open_credits),
            "notes": f"Crédito unificado de {c_name} (Plazo {p_days} días)",
            "payments_count": sum(len(c.payments or []) for c in p_credits)
        })

    return items


# ===== 3. GET SINGLE CREDIT DETAIL =====
@router.get("/invoices/{invoice_id}")
async def get_customer_invoice_credit(
    invoice_id: str,
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Obtener detalle unificado de cuenta por cobrar de cliente con sus facturas y abonos."""
    try:
        target_uuid = uuid.UUID(invoice_id)
    except ValueError:
        raise HTTPException(status_code=400, detail="ID inválido")

    c_uuid = uuid.UUID(str(current_user.company_id))
    await _sync_sales_to_credits(db, c_uuid)

    # Buscar cliente o factura
    p_stmt = select(Partner).where(Partner.id == target_uuid, Partner.company_id == c_uuid)
    p_res = await db.execute(p_stmt)
    partner = p_res.scalar_one_or_none()

    if not partner:
        i_stmt = select(CustomerInvoiceCredit).where(CustomerInvoiceCredit.id == target_uuid, CustomerInvoiceCredit.company_id == c_uuid)
        i_res = await db.execute(i_stmt)
        inv = i_res.scalar_one_or_none()
        if inv and inv.customer_id:
            p_res = await db.execute(select(Partner).where(Partner.id == inv.customer_id, Partner.company_id == c_uuid))
            partner = p_res.scalar_one_or_none()

    if not partner:
        raise HTTPException(status_code=404, detail="Cliente o cuenta no encontrada")

    # Obtener historial de facturas a crédito y abonos
    c_stmt = (
        select(CustomerInvoiceCredit)
        .options(selectinload(CustomerInvoiceCredit.payments))
        .where(
            CustomerInvoiceCredit.customer_id == partner.id,
            CustomerInvoiceCredit.company_id == c_uuid,
            CustomerInvoiceCredit.is_deleted == False
        )
        .order_by(CustomerInvoiceCredit.issue_date.desc())
    )
    c_res = await db.execute(c_stmt)
    credits = c_res.scalars().all()

    p_stmt = (
        select(CustomerCreditPayment)
        .where(
            CustomerCreditPayment.customer_id == partner.id,
            CustomerCreditPayment.company_id == c_uuid,
            CustomerCreditPayment.is_deleted == False
        )
        .order_by(CustomerCreditPayment.payment_date.desc())
    )
    p_res = await db.execute(p_stmt)
    payments = p_res.scalars().all()

    total_invoiced = sum(float(c.total or 0.0) for c in credits) if credits else float(partner.balance or 0.0)
    open_credits_bal = sum(float(c.balance or 0.0) for c in credits if (c.balance or 0) > 0)
    curr_balance = max(float(partner.balance or 0.0), open_credits_bal)

    if float(partner.balance or 0.0) < open_credits_bal:
        partner.balance = Decimal(str(open_credits_bal))
        await db.commit()
    
    first_dt = credits[-1].issue_date if credits else partner.created_at
    due_dt = (first_dt + timedelta(days=partner.payment_days or 30)) if first_dt else None

    return {
        "id": str(partner.id),
        "number": f"CLI-{str(partner.id)[:6].upper()}",
        "customer_id": str(partner.id),
        "customer_name": partner.name,
        "customer_nit": partner.nit or "CF",
        "customer_phone": partner.phone or partner.mobile or "",
        "issue_date": first_dt.strftime("%d/%m/%Y") if first_dt else "",
        "due_date": due_dt.strftime("%d/%m/%Y") if due_dt else "Sin fecha",
        "payment_days": partner.payment_days or 30,
        "total": max(total_invoiced, curr_balance),
        "balance": curr_balance,
        "credit_limit": float(partner.credit_limit or 0.0),
        "state": "open" if curr_balance > 0 else "paid",
        "notes": f"Crédito unificado de {partner.name}",
        "credits": [
            {
                "id": str(c.id),
                "number": c.number,
                "date": c.issue_date.strftime("%d/%m/%Y") if c.issue_date else "",
                "total": float(c.total or 0.0),
                "balance": float(c.balance or 0.0),
                "notes": c.notes or ""
            }
            for c in credits
        ],
        "payments": [
            {
                "id": str(pmt.id),
                "receipt_number": pmt.receipt_number,
                "amount": float(pmt.amount),
                "date": pmt.payment_date.strftime("%d/%m/%Y %H:%M") if pmt.payment_date else "",
                "method": pmt.payment_method,
                "reference": pmt.reference or "",
                "notes": pmt.notes or ""
            }
            for pmt in payments
        ]
    }


# ===== 4. REGISTER NEW CREDIT ACCOUNT =====
@router.post("/invoices")
async def create_customer_invoice_credit(
    data: CustomerInvoiceCreditCreateIn,
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Crear manualmente una cuenta de crédito a cliente."""
    c_uuid = uuid.UUID(str(current_user.company_id))

    cust_uuid = None
    if data.customer_id:
        try:
            cust_uuid = uuid.UUID(data.customer_id)
        except ValueError:
            pass

    partner = None
    if cust_uuid:
        p_res = await db.execute(select(Partner).where(Partner.id == cust_uuid, Partner.company_id == c_uuid))
        partner = p_res.scalar_one_or_none()

    if not partner and data.customer_name:
        partner = Partner(
            company_id=c_uuid,
            name=data.customer_name,
            nit=data.customer_nit or "CF",
            credit_limit=Decimal("5000.00"),
            payment_days=data.credit_days or 30,
            partner_type=PartnerType.CUSTOMER
        )
        db.add(partner)
        await db.flush()
        cust_uuid = partner.id

    now = datetime.now(timezone.utc)
    i_date = now
    if data.issue_date:
        try:
            i_date = datetime.strptime(data.issue_date, "%Y-%m-%d").replace(tzinfo=timezone.utc)
        except Exception:
            pass

    due_date = i_date + timedelta(days=data.credit_days or 30)

    credit_inv = CustomerInvoiceCredit(
        company_id=c_uuid,
        customer_id=cust_uuid,
        number=data.number,
        issue_date=i_date,
        due_date=due_date,
        total=Decimal(str(data.total)),
        balance=Decimal(str(data.total)),
        state="open",
        notes=data.notes
    )
    db.add(credit_inv)

    if partner:
        partner.balance = (partner.balance or Decimal("0")) + Decimal(str(data.total))

    await db.commit()
    await db.refresh(credit_inv)

    return {
        "id": str(credit_inv.id),
        "number": credit_inv.number,
        "total": float(credit_inv.total),
        "balance": float(credit_inv.balance),
        "due_date": credit_inv.due_date.strftime("%d/%m/%Y") if credit_inv.due_date else ""
    }


# ===== 5. REGISTER PAYMENT / ABONO AL CLIENTE =====
@router.post("/invoices/{invoice_id}/payments")
async def register_customer_payment(
    invoice_id: str,
    data: ARPaymentIn,
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Registrar un abono al cliente (aplicando a su saldo y a sus facturas de crédito FIFO)."""
    check_credit_payment_permission(current_user)
    try:
        target_uuid = uuid.UUID(invoice_id)
    except ValueError:
        raise HTTPException(status_code=400, detail="ID inválido")

    c_uuid = uuid.UUID(str(current_user.company_id))

    # Buscar si es Partner o Invoice
    p_stmt = select(Partner).where(Partner.id == target_uuid, Partner.company_id == c_uuid)
    p_res = await db.execute(p_stmt)
    partner = p_res.scalar_one_or_none()

    inv = None
    if not partner:
        i_stmt = select(CustomerInvoiceCredit).where(CustomerInvoiceCredit.id == target_uuid, CustomerInvoiceCredit.company_id == c_uuid)
        i_res = await db.execute(i_stmt)
        inv = i_res.scalar_one_or_none()
        if inv and inv.customer_id:
            p_res = await db.execute(select(Partner).where(Partner.id == inv.customer_id, Partner.company_id == c_uuid))
            partner = p_res.scalar_one_or_none()

    if not partner and not inv:
        raise HTTPException(status_code=404, detail="Cliente o cuenta por cobrar no encontrada")

    cust_id = partner.id if partner else inv.customer_id
    total_balance = partner.balance if partner else inv.balance

    amt = Decimal(str(data.amount))
    if amt <= 0:
        raise HTTPException(status_code=400, detail="El monto del abono debe ser mayor a 0.")

    if total_balance and amt > total_balance:
        amt = total_balance

    # Correlativo de recibo
    p_count_stmt = select(func.count(CustomerCreditPayment.id)).where(CustomerCreditPayment.company_id == c_uuid)
    p_count_res = await db.execute(p_count_stmt)
    receipt_num = f"REC-{(p_count_res.scalar() or 0) + 1:06d}"

    # Aplicar a las facturas de crédito abiertas de este cliente en orden FIFO
    inv_list_stmt = (
        select(CustomerInvoiceCredit)
        .where(
            CustomerInvoiceCredit.customer_id == cust_id,
            CustomerInvoiceCredit.company_id == c_uuid,
            CustomerInvoiceCredit.balance > 0,
            CustomerInvoiceCredit.is_deleted == False
        )
        .order_by(CustomerInvoiceCredit.issue_date.asc())
    )
    inv_list_res = await db.execute(inv_list_stmt)
    open_invs = inv_list_res.scalars().all()

    remaining_to_apply = amt
    primary_inv_id = None
    for open_inv in open_invs:
        if remaining_to_apply <= 0:
            break
        if not primary_inv_id:
            primary_inv_id = open_inv.id
        to_pay = min(open_inv.balance, remaining_to_apply)
        open_inv.balance = open_inv.balance - to_pay
        if open_inv.balance <= 0:
            open_inv.balance = Decimal("0")
            open_inv.state = "paid"
        remaining_to_apply -= to_pay

    payment = CustomerCreditPayment(
        company_id=c_uuid,
        invoice_id=primary_inv_id or (inv.id if inv else None),
        customer_id=cust_id,
        receipt_number=receipt_num,
        amount=amt,
        payment_date=datetime.now(timezone.utc),
        payment_method=data.payment_method,
        reference=data.reference,
        notes=data.notes
    )
    db.add(payment)

    # Actualizar saldo del partner
    prev_balance = partner.balance if partner else amt
    if partner and partner.balance is not None:
        partner.balance = max(Decimal("0"), partner.balance - amt)

    await db.commit()
    await db.refresh(payment)

    new_bal = float(partner.balance) if partner else 0.0

    return {
        "status": "success",
        "payment_id": str(payment.id),
        "receipt_number": payment.receipt_number,
        "previous_balance": float(prev_balance),
        "paid_amount": float(amt),
        "new_balance": new_bal,
        "message": "¡Abono registrado exitosamente!"
    }


def numero_a_letras_quetzales(numero: float) -> str:
    """Convierte un importe numérico a su representación en letras en Quetzales."""
    try:
        entero = int(numero)
        centavos = int(round((numero - entero) * 100))
        
        unidades = ["", "UN", "DOS", "TRES", "CUATRO", "CINCO", "SEIS", "SIETE", "OCHO", "NUEVE"]
        especiales = ["DIEZ", "ONCE", "DOCE", "TRECE", "CATORCE", "QUINCE", "DIECISEIS", "DIECISIETE", "DIECIOCHO", "DIECINUEVE", "VEINTE"]
        veintes = ["VEINTIUN", "VEINTIDOS", "VEINTITRES", "VEINTICUATRO", "VEINTICINCO", "VEINTISEIS", "VEINTISIETE", "VEINTIOCHO", "VEINTINUEVE"]
        decenas = ["", "DIEZ", "VEINTE", "TREINTA", "CUARENTA", "CINCUENTA", "SESENTA", "SETENTA", "OCHENTA", "NOVENTA"]
        centenas = ["", "CIENTO", "DOSCIENTOS", "TRESCIENTOS", "CUATROCIENTOS", "QUINIENTOS", "SEISCIENTOS", "SETECIENTOS", "OCHOCIENTOS", "NOVECIENTOS"]

        def _centenas(n):
            if n == 100:
                return "CIEN"
            c = n // 100
            d = (n % 100) // 10
            u = n % 10
            txt = centenas[c]
            resto = n % 100
            if resto == 0:
                return txt
            if txt:
                txt += " "
            if resto <= 20:
                if resto <= 9:
                    txt += unidades[resto]
                elif resto == 10:
                    txt += "DIEZ"
                elif resto <= 20:
                    txt += especiales[resto - 10]
            elif resto < 30:
                txt += veintes[resto - 21]
            else:
                txt += decenas[d]
                if u > 0:
                    txt += " Y " + unidades[u]
            return txt.strip()

        def _miles(n):
            mil = n // 1000
            resto = n % 1000
            if mil == 0:
                return _centenas(resto)
            if mil == 1:
                txt = "MIL"
            else:
                txt = _centenas(mil) + " MIL"
            if resto > 0:
                txt += " " + _centenas(resto)
            return txt.strip()

        if entero == 0:
            letras = "CERO"
        elif entero < 1000000:
            letras = _miles(entero)
        else:
            millon = entero // 1000000
            resto = entero % 1000000
            txt_millon = "UN MILLON" if millon == 1 else _centenas(millon) + " MILLONES"
            if resto > 0:
                letras = txt_millon + " " + _miles(resto)
            else:
                letras = txt_millon

        moneda = "QUETZAL" if entero == 1 else "QUETZALES"
        if centavos > 0:
            return f"{letras} {moneda} CON {centavos:02d}/100"
        else:
            return f"{letras} {moneda} EXACTOS"
    except Exception:
        return f"{numero:.2f} QUETZALES"


def fecha_en_letras_espanol(dt: datetime) -> str:
    """Retorna fecha en formato largo 'JUEVES 22 DE OCTUBRE DE 2026'."""
    dias = ["LUNES", "MARTES", "MIÉRCOLES", "JUEVES", "VIERNES", "SÁBADO", "DOMINGO"]
    meses = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]
    try:
        dia_nom = dias[dt.weekday()]
        mes_nom = meses[dt.month - 1]
        return f"{dia_nom} {dt.day} DE {mes_nom} DE {dt.year}"
    except Exception:
        return dt.strftime("%d/%m/%Y")


# ===== 6. PAYMENT RECEIPT / COMPROBANTE DE ABONO =====
@router.get("/payments/{payment_id}/receipt")
async def get_payment_receipt(
    payment_id: str,
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Obtener datos para recibo de caja / abono imprimible para el cliente en formato Media Carta."""
    try:
        p_uuid = uuid.UUID(payment_id)
    except ValueError:
        raise HTTPException(status_code=400, detail="ID de pago inválido")

    c_uuid = uuid.UUID(str(current_user.company_id))

    stmt = (
        select(CustomerCreditPayment)
        .options(
            selectinload(CustomerCreditPayment.invoice),
            selectinload(CustomerCreditPayment.customer)
        )
        .where(
            CustomerCreditPayment.id == p_uuid,
            CustomerCreditPayment.company_id == c_uuid
        )
    )
    res = await db.execute(stmt)
    pmt = res.scalar_one_or_none()
    if not pmt:
        raise HTTPException(status_code=404, detail="Comprobante de abono no encontrado")

    inv = pmt.invoice
    if not inv and pmt.invoice_id:
        i_res = await db.execute(select(CustomerInvoiceCredit).where(CustomerInvoiceCredit.id == pmt.invoice_id))
        inv = i_res.scalar_one_or_none()

    customer = pmt.customer
    if not customer and inv and inv.customer_id:
        c_res = await db.execute(select(Partner).where(Partner.id == inv.customer_id))
        customer = c_res.scalar_one_or_none()
    elif not customer and pmt.customer_id:
        c_res = await db.execute(select(Partner).where(Partner.id == pmt.customer_id))
        customer = c_res.scalar_one_or_none()

    comp_res = await db.execute(select(Company).where(Company.id == c_uuid))
    company = comp_res.scalar_one_or_none()

    dt = pmt.payment_date or datetime.now(timezone.utc)
    amt = float(pmt.amount)

    concept = pmt.notes
    if not concept:
        if inv and inv.number:
            concept = f"Abono a cuenta por cobrar Factura / Venta #{inv.number}"
        else:
            concept = "Abono a cuenta por cobrar de crédito comercial"

    return {
        "receipt_id": str(pmt.id),
        "receipt_number": pmt.receipt_number,
        "payment_date": dt.strftime("%d/%m/%Y %H:%M"),
        "date_short": dt.strftime("%d/%m/%Y"),
        "date_long": fecha_en_letras_espanol(dt),
        "payment_method": pmt.payment_method,
        "reference": pmt.reference or "N/A",
        "notes": pmt.notes or "",
        "concept": concept,
        "amount_paid": amt,
        "amount_letters": numero_a_letras_quetzales(amt),
        "invoice": {
            "id": str(inv.id) if inv else "",
            "number": inv.number if inv else "Cobro Directo",
            "total": float(inv.total) if inv else amt,
            "current_balance": float(inv.balance) if inv else 0.0,
            "previous_balance": float((inv.balance or Decimal("0")) + pmt.amount) if inv else amt,
            "issue_date": inv.issue_date.strftime("%d/%m/%Y") if inv and inv.issue_date else "",
            "due_date": inv.due_date.strftime("%d/%m/%Y") if inv and inv.due_date else ""
        },
        "customer": {
            "id": str(customer.id) if customer else "",
            "name": customer.name if customer else "Consumidor Final",
            "nit": customer.nit if customer else "CF",
            "phone": customer.phone or customer.mobile if customer else "",
            "email": customer.email if customer else "",
            "address": customer.address if customer else ""
        },
        "company": {
            "name": company.name if company else "Productos Automotrices del Sur S.A.",
            "title_name": "Productos Automotrices del Sur S.A.",
            "nit": company.nit if company else "11563009-0",
            "address": "Lote 67, Calle Central, Nueva Concepción, Escuintla.",
            "emails": "proausur78@gmail.com | proausu.sa@gmail.com",
            "phone": company.phone if company else "+(502) 3454-3857"
        },
        "user_name": getattr(current_user, 'full_name', '') or getattr(current_user, 'username', '') or getattr(current_user, 'email', 'Operador')
    }


# ===== 7. AGING REPORT (ANTIGÜEDAD DE SALDOS) =====
@router.get("/aging")
async def get_ar_aging_report(
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Reporte ejecutivo de Antigüedad de Saldos por Cliente (0-30, 31-60, 61-90, 90+)."""
    c_uuid = uuid.UUID(str(current_user.company_id))
    await _sync_sales_to_credits(db, c_uuid)
    now = datetime.now(timezone.utc)

    stmt = (
        select(CustomerInvoiceCredit)
        .options(selectinload(CustomerInvoiceCredit.customer))
        .where(
            CustomerInvoiceCredit.company_id == c_uuid,
            CustomerInvoiceCredit.balance > 0,
            CustomerInvoiceCredit.is_deleted == False
        )
    )
    res = await db.execute(stmt)
    invoices = res.scalars().all()

    summary = {
        "current": 0.0,
        "range_1_30": 0.0,
        "range_31_60": 0.0,
        "range_61_90": 0.0,
        "range_90_plus": 0.0,
        "total": 0.0
    }

    customers_dict = {}

    for inv in invoices:
        bal = float(inv.balance or 0.0)
        due_dt = inv.due_date
        if due_dt:
            if due_dt.tzinfo is None:
                due_dt = due_dt.replace(tzinfo=timezone.utc)
            days_overdue = (now.date() - due_dt.date()).days
        else:
            days_overdue = 0

        cust_id = str(inv.customer_id) if inv.customer_id else "unknown"
        cust_name = inv.customer.name if inv.customer else "Cliente No Identificado"
        cust_nit = inv.customer.nit if inv.customer else "CF"
        cust_limit = float(inv.customer.credit_limit or 0.0) if inv.customer else 0.0

        if cust_id not in customers_dict:
            customers_dict[cust_id] = {
                "customer_id": cust_id,
                "customer_name": cust_name,
                "customer_nit": cust_nit,
                "credit_limit": cust_limit,
                "current": 0.0,
                "range_1_30": 0.0,
                "range_31_60": 0.0,
                "range_61_90": 0.0,
                "range_90_plus": 0.0,
                "total": 0.0,
                "invoices_count": 0
            }

        target = customers_dict[cust_id]
        target["invoices_count"] += 1
        target["total"] += bal
        summary["total"] += bal

        if days_overdue <= 0:
            target["current"] += bal
            summary["current"] += bal
        elif 1 <= days_overdue <= 30:
            target["range_1_30"] += bal
            summary["range_1_30"] += bal
        elif 31 <= days_overdue <= 60:
            target["range_31_60"] += bal
            summary["range_31_60"] += bal
        elif 61 <= days_overdue <= 90:
            target["range_61_90"] += bal
            summary["range_61_90"] += bal
        else:
            target["range_90_plus"] += bal
            summary["range_90_plus"] += bal

    customer_list = list(customers_dict.values())
    customer_list.sort(key=lambda x: x["total"], reverse=True)

    return {
        "summary": summary,
        "customers": customer_list
    }


# ===== 8. CUSTOMER STATEMENT (ESTADO DE CUENTA 360) =====
@router.get("/customers/{customer_id}/statement")
async def get_customer_account_statement(
    customer_id: str,
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Estado de cuenta individual de un cliente con saldo, límite y movimientos."""
    c_uuid = uuid.UUID(str(current_user.company_id))
    await _sync_sales_to_credits(db, c_uuid)

    try:
        cust_uuid = uuid.UUID(customer_id)
    except ValueError:
        raise HTTPException(status_code=400, detail="ID de cliente inválido")

    c_stmt = select(Partner).where(
        Partner.id == cust_uuid,
        Partner.company_id == c_uuid
    )
    c_res = await db.execute(c_stmt)
    customer = c_res.scalar_one_or_none()
    if not customer:
        raise HTTPException(status_code=404, detail="Cliente no encontrado")

    # Invoices
    inv_stmt = (
        select(CustomerInvoiceCredit)
        .options(selectinload(CustomerInvoiceCredit.payments))
        .where(
            CustomerInvoiceCredit.customer_id == cust_uuid,
            CustomerInvoiceCredit.company_id == current_user.company_id,
            CustomerInvoiceCredit.is_deleted == False
        )
        .order_by(CustomerInvoiceCredit.issue_date.desc())
    )
    inv_res = await db.execute(inv_stmt)
    invoices = inv_res.scalars().all()

    # Company info
    comp_stmt = select(Company).where(Company.id == c_uuid)
    comp_res = await db.execute(comp_stmt)
    company = comp_res.scalar_one_or_none()

    now_dt = datetime.now()
    today_date = now_dt.date()
    dias_semana = ['lunes', 'martes', 'miércoles', 'jueves', 'viernes', 'sábado', 'domingo']
    meses_nombre = ['enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre']
    date_long = f"{dias_semana[now_dt.weekday()]}, {now_dt.day} de {meses_nombre[now_dt.month - 1]} de {now_dt.year}"
    date_footer = f"{now_dt.day} de {meses_nombre[now_dt.month - 1]} de {now_dt.year}"
    time_str = now_dt.strftime("%H:%M")

    # Sort invoices by issue_date asc
    invoices_sorted = sorted(invoices, key=lambda x: x.issue_date or datetime.min)

    total_envio = 0.0
    total_debt = 0.0
    total_vencido = 0.0
    total_abono = 0.0
    invoices_list = []

    for idx, i in enumerate(invoices_sorted, start=1):
        inv_tot = float(i.total or 0)
        inv_bal = float(i.balance or 0)
        inv_abono = max(0.0, inv_tot - inv_bal)

        if i.due_date:
            due_d = i.due_date.date() if hasattr(i.due_date, 'date') else i.due_date
            days_diff = (due_d - today_date).days
        else:
            days_diff = 0

        status_deuda = "Vencido" if days_diff < 0 else "No vencido"

        total_envio += inv_tot
        total_debt += inv_bal
        total_abono += inv_abono
        if days_diff < 0 and inv_bal > 0.001:
            total_vencido += inv_bal

        invoices_list.append({
            "no": idx,
            "id": str(i.id),
            "customer_code": str(customer.code) if customer.code else str(customer.id)[:8],
            "number": i.number,
            "issue_date": i.issue_date.strftime("%d/%m/%Y") if i.issue_date else "",
            "due_date": i.due_date.strftime("%d/%m/%Y") if i.due_date else "",
            "issue_date_short": i.issue_date.strftime("%d-%m-%y") if i.issue_date else "",
            "due_date_short": i.due_date.strftime("%d-%m-%y") if i.due_date else "",
            "total": inv_tot,
            "balance": inv_bal,
            "abono": inv_abono,
            "days": days_diff,
            "status_deuda": status_deuda,
            "state": i.state
        })

    limit = float(customer.credit_limit or 0.0)
    available_credit = max(0.0, limit - total_debt)

    # Payments
    pay_stmt = (
        select(CustomerCreditPayment)
        .where(
            CustomerCreditPayment.customer_id == cust_uuid,
            CustomerCreditPayment.company_id == current_user.company_id
        )
        .order_by(CustomerCreditPayment.payment_date.desc())
    )
    pay_res = await db.execute(pay_stmt)
    payments = pay_res.scalars().all()

    return {
        "customer": {
            "id": str(customer.id),
            "code": str(customer.code) if customer.code else str(customer.id)[:8],
            "name": customer.name,
            "nit": customer.nit or "C/F",
            "phone": customer.phone or customer.mobile or "",
            "address": customer.address or "",
            "email": customer.email or "",
            "credit_limit": limit,
            "payment_days": customer.payment_days or 30,
            "total_debt": round(total_debt, 2),
            "available_credit": round(available_credit, 2),
            "credit_status": str(customer.credit_status.value if hasattr(customer.credit_status, 'value') else customer.credit_status)
        },
        "company": {
            "name": company.name if company and company.name else "Productos Automotrices del Sur, S.A.",
            "title_name": "Productos Automotrices del Sur, S.A.",
            "nit": company.nit if company and company.nit else "11563009-0",
            "address": company.address if company and company.address else "Lote 67, Calle Central, Nueva Concepción, Esc.",
            "phone": company.phone if company and company.phone else "3454-3857",
            "email": company.email if company and company.email else "proausu.sa@gmail.com"
        },
        "summary": {
            "total_envio": round(total_envio, 2),
            "total_deuda": round(total_debt, 2),
            "total_vencido": round(total_vencido, 2),
            "total_abono": round(total_abono, 2)
        },
        "date_long": date_long,
        "date_footer": date_footer,
        "time_str": time_str,
        "invoices": invoices_list,
        "payments": [
            {
                "id": str(p.id),
                "receipt_number": p.receipt_number,
                "amount": float(p.amount),
                "payment_date": p.payment_date.strftime("%d/%m/%Y %H:%M") if p.payment_date else "",
                "method": p.payment_method,
                "reference": p.reference or ""
            }
            for p in payments
        ]
    }


# ===== 9. EXPORT CSV =====
@router.get("/export")
async def export_ar_csv(
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Descargar estado general de Cuentas por Cobrar en formato CSV."""
    c_uuid = uuid.UUID(str(current_user.company_id))
    await _sync_sales_to_credits(db, c_uuid)
    now = datetime.now(timezone.utc)

    stmt = (
        select(CustomerInvoiceCredit)
        .options(selectinload(CustomerInvoiceCredit.customer))
        .where(
            CustomerInvoiceCredit.company_id == c_uuid,
            CustomerInvoiceCredit.is_deleted == False
        )
        .order_by(CustomerInvoiceCredit.due_date.asc())
    )
    res = await db.execute(stmt)
    credits = res.scalars().all()

    output = io.StringIO()
    output.write('\ufeff')
    writer = csv.writer(output)

    writer.writerow([
        "Cliente",
        "NIT",
        "No. Factura / Credito",
        "Fecha Emision",
        "Fecha Vencimiento",
        "Dias Mora / Restantes",
        "Monto Total (Q)",
        "Saldo Pendiente (Q)",
        "Estado",
        "Notas"
    ])

    for c in credits:
        bal = float(c.balance or 0.0)
        due_dt = c.due_date
        days_diff = 0
        status_txt = "PAGADA"
        if bal > 0:
            if due_dt:
                if due_dt.tzinfo is None:
                    due_dt = due_dt.replace(tzinfo=timezone.utc)
                diff = (due_dt.date() - now.date()).days
                days_diff = diff
                if diff < 0:
                    status_txt = f"VENCIDA ({abs(diff)} dias)"
                elif diff <= 7:
                    status_txt = f"POR VENCER ({diff} dias)"
                else:
                    status_txt = f"AL DIA ({diff} dias)"
            else:
                status_txt = "AL DIA"

        writer.writerow([
            c.customer.name if c.customer else "N/A",
            c.customer.nit if c.customer else "CF",
            c.number,
            c.issue_date.strftime("%d/%m/%Y") if c.issue_date else "",
            c.due_date.strftime("%d/%m/%Y") if c.due_date else "",
            days_diff,
            f"{float(c.total or 0.0):.2f}",
            f"{bal:.2f}",
            status_txt,
            c.notes or ""
        ])

    csv_data = output.getvalue()
    filename = f"cuentas_por_cobrar_{now.strftime('%Y%m%d_%H%M')}.csv"

    return Response(
        content=csv_data,
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


# ===== 10. CUSTOMERS LOOKUP FOR AUTOCOMPLETE =====
@router.get("/customers-lookup")
async def lookup_ar_customers(
    q: Optional[str] = Query(None),
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Búsqueda rápida de clientes para autocompletado en CxC y abonos."""
    c_uuid = uuid.UUID(str(current_user.company_id))
    await _sync_sales_to_credits(db, c_uuid)

    stmt = select(Partner).where(
        Partner.company_id == c_uuid,
        Partner.is_deleted == False
    ).order_by(Partner.name.asc()).limit(50)

    if q:
        stmt = stmt.where(
            (Partner.name.ilike(f"%{q}%")) | (Partner.nit.ilike(f"%{q}%"))
        )

    res = await db.execute(stmt)
    partners = res.scalars().all()

    items = []
    for p in partners:
        limit = float(p.credit_limit or 0.0)
        bal = float(p.balance or 0.0)
        items.append({
            "id": str(p.id),
            "name": p.name,
            "nit": p.nit or "CF",
            "phone": p.phone or p.mobile or "",
            "email": p.email or "",
            "credit_limit": limit,
            "balance": bal,
            "available_credit": max(0.0, limit - bal),
            "payment_days": p.payment_days or 30,
            "credit_status": str(p.credit_status.value if hasattr(p.credit_status, 'value') else p.credit_status)
        })
    return {"items": items}


# ===== 11. EXPORT CUSTOMER STATEMENT CSV =====
@router.get("/customers/{customer_id}/export")
async def export_customer_statement_csv(
    customer_id: str,
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Descargar estado de cuenta individual de un cliente en formato CSV."""
    cust_uuid = uuid.UUID(customer_id)
    c_stmt = select(Partner).where(
        Partner.id == cust_uuid,
        Partner.company_id == current_user.company_id
    )
    c_res = await db.execute(c_stmt)
    customer = c_res.scalar_one_or_none()
    if not customer:
        raise HTTPException(status_code=404, detail="Cliente no encontrado")

    inv_stmt = select(CustomerInvoiceCredit).where(
        CustomerInvoiceCredit.customer_id == cust_uuid,
        CustomerInvoiceCredit.company_id == current_user.company_id,
        CustomerInvoiceCredit.is_deleted == False
    ).order_by(CustomerInvoiceCredit.issue_date.asc())
    inv_res = await db.execute(inv_stmt)
    invoices = inv_res.scalars().all()

    pay_stmt = select(CustomerCreditPayment).where(
        CustomerCreditPayment.customer_id == cust_uuid,
        CustomerCreditPayment.company_id == current_user.company_id
    ).order_by(CustomerCreditPayment.payment_date.asc())
    pay_res = await db.execute(pay_stmt)
    payments = pay_res.scalars().all()

    output = io.StringIO()
    output.write('\ufeff')
    writer = csv.writer(output)

    writer.writerow(["ESTADO DE CUENTA DE CLIENTE - POS EMPRESARIAL"])
    writer.writerow(["Cliente", customer.name, "NIT", customer.nit])
    writer.writerow(["Limite de Credito (Q)", f"{float(customer.credit_limit or 0):.2f}", "Dias de Credito", customer.payment_days or 30])
    writer.writerow([])
    writer.writerow(["FACTURAS Y CARGOS A CREDITO"])
    writer.writerow(["No. Documento", "Fecha Emision", "Fecha Vencimiento", "Total Facturado (Q)", "Saldo Pendiente (Q)", "Estado"])
    for i in invoices:
        writer.writerow([
            i.number,
            i.issue_date.strftime("%d/%m/%Y") if i.issue_date else "",
            i.due_date.strftime("%d/%m/%Y") if i.due_date else "",
            f"{float(i.total or 0):.2f}",
            f"{float(i.balance or 0):.2f}",
            i.state.upper()
        ])

    writer.writerow([])
    writer.writerow(["HISTORIAL DE ABONOS Y PAGOS RECIBIDOS"])
    writer.writerow(["No. Recibo", "Fecha Abono", "Metodo de Pago", "Referencia", "Monto Abonado (Q)"])
    for p in payments:
        writer.writerow([
            p.receipt_number,
            p.payment_date.strftime("%d/%m/%Y %H:%M") if p.payment_date else "",
            p.payment_method.upper(),
            p.reference or "N/A",
            f"{float(p.amount or 0):.2f}"
        ])

    csv_data = output.getvalue()
    clean_name = "".join(c for c in customer.name if c.isalnum() or c in (' ', '_', '-')).strip().replace(' ', '_')
    filename = f"estado_cuenta_{clean_name}_{datetime.now().strftime('%Y%m%d')}.csv"

    return Response(
        content=csv_data,
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


# ===== 12. EXPORT CUSTOMER STATEMENT EXCEL (.XLSX) OFICIAL =====
def _build_statement_excel(customer_dict, invoices_list, company_dict, summary_dict, date_long, date_footer, time_str) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Estado de Cuenta"
    ws.views.sheetView[0].showGridLines = True

    # Column widths matching official document layout
    col_widths = {
        'A': 6,   # No.
        'B': 13,  # ID Cliente
        'C': 16,  # Factura No.
        'D': 15,  # Fecha Factura
        'E': 15,  # Fecha Vencim.
        'F': 18,  # Total de Factura
        'G': 16,  # Status Deuda
        'H': 12,  # # de Días
        'I': 16,  # Abono
        'J': 18   # Saldo Actual
    }
    for col, w in col_widths.items():
        ws.column_dimensions[col].width = w

    thin_border_side = Side(style='thin', color='B0BEC5')
    dotted_border_side = Side(style='dotted', color='94A3B8')

    # --- LEFT CARD: Company & Logos (A2:F11) ---
    for r in range(2, 12):
        for c in range(1, 7):
            cell = ws.cell(row=r, column=c)
            l = thin_border_side if c == 1 else None
            ri = thin_border_side if c == 6 else None
            t = thin_border_side if r == 2 else None
            b = thin_border_side if r == 11 else None
            cell.border = Border(left=l, right=ri, top=t, bottom=b)

    # Banner image
    banner_path = os.path.join(os.getcwd(), 'frontend', 'static', 'img', 'statement_logos_banner.png')
    if os.path.exists(banner_path):
        try:
            img = OpenpyxlImage(banner_path)
            img.width = 240
            img.height = 49
            ws.add_image(img, 'A2')
        except Exception:
            pass

    ws.cell(row=6, column=1, value=company_dict.get('name', 'Productos Automotrices del Sur, S.A.')).font = Font(name='Arial', size=11, bold=True, color='0F172A')

    ws.cell(row=7, column=1, value='NIT:').font = Font(name='Arial', size=9, bold=True, color='1E293B')
    ws.cell(row=7, column=2, value=company_dict.get('nit', '11563009-0')).font = Font(name='Arial', size=9, color='334155')

    ws.cell(row=8, column=1, value='Dirección:').font = Font(name='Arial', size=9, bold=True, color='1E293B')
    ws.cell(row=8, column=2, value=company_dict.get('address', 'Lote 67, Calle Central, Nueva Concepción, Esc.')).font = Font(name='Arial', size=9, color='334155')

    ws.cell(row=9, column=1, value='Teléfono:').font = Font(name='Arial', size=9, bold=True, color='1E293B')
    ws.cell(row=9, column=2, value=company_dict.get('phone', '3454-3857')).font = Font(name='Arial', size=9, color='334155')

    ws.cell(row=10, column=1, value='Gmail:').font = Font(name='Arial', size=9, bold=True, color='1E293B')
    ws.cell(row=10, column=2, value=company_dict.get('email', 'proausu.sa@gmail.com')).font = Font(name='Arial', size=9, color='334155')

    # --- RIGHT CARD: Estado de Cuenta (G2:J11) ---
    for r in range(2, 12):
        for c in range(7, 11):
            cell = ws.cell(row=r, column=c)
            l = thin_border_side if c == 7 else None
            ri = thin_border_side if c == 10 else None
            t = thin_border_side if r == 2 else None
            b = thin_border_side if r == 11 else None
            cell.border = Border(left=l, right=ri, top=t, bottom=b)

    ws.merge_cells('G2:J3')
    st_cell = ws.cell(row=2, column=7, value='Estado de Cuenta')
    st_cell.font = Font(name='Arial', size=13, bold=True, color='0F172A')
    st_cell.alignment = Alignment(horizontal='center', vertical='center')
    header_fill = PatternFill(start_color='EBF1F6', end_color='EBF1F6', fill_type='solid')
    for r in (2, 3):
        for c in range(7, 11):
            ws.cell(row=r, column=c).fill = header_fill
            if r == 3:
                ws.cell(row=r, column=c).border = Border(bottom=thin_border_side, left=thin_border_side if c == 7 else None, right=thin_border_side if c == 10 else None)

    def set_summary_row(row, label, val_num=None, val_str=None):
        ws.merge_cells(start_row=row, start_column=7, end_row=row, end_column=8)
        lbl_c = ws.cell(row=row, column=7, value=label)
        lbl_c.font = Font(name='Arial', size=10, bold=True, color='1E293B')
        lbl_c.alignment = Alignment(horizontal='left', vertical='center')

        q_c = ws.cell(row=row, column=9, value='Q')
        q_c.font = Font(name='Arial', size=10, bold=True, color='1E293B')
        q_c.alignment = Alignment(horizontal='right', vertical='center')

        v_c = ws.cell(row=row, column=10)
        if val_num is not None:
            v_c.value = val_num
            v_c.number_format = '#,##0.00'
        else:
            v_c.value = val_str or '-'
        v_c.font = Font(name='Arial', size=10, bold=True, color='1E293B')
        v_c.alignment = Alignment(horizontal='right', vertical='center')

    set_summary_row(5, 'Total Envío:', val_num=summary_dict.get('total_envio', 0.0))
    set_summary_row(6, 'Total Deuda:', val_num=summary_dict.get('total_deuda', 0.0))
    set_summary_row(7, 'Total Vencido:', val_num=summary_dict.get('total_vencido', 0.0))
    tot_abono = summary_dict.get('total_abono', 0.0)
    set_summary_row(8, 'Total de Abono:', val_num=tot_abono if tot_abono > 0 else None, val_str='-' if tot_abono <= 0 else None)

    # --- CLIENT INFO CARD (A13:J17) ---
    for r in range(13, 18):
        for c in range(1, 11):
            cell = ws.cell(row=r, column=c)
            l = thin_border_side if c == 1 else None
            ri = thin_border_side if c == 10 else None
            t = thin_border_side if r == 13 else None
            b = thin_border_side if r == 17 else None
            cell.border = Border(left=l, right=ri, top=t, bottom=b)

    # Row 13: ID Cliente & Fecha
    ws.cell(row=13, column=1, value='ID Cliente:').font = Font(name='Arial', size=10, bold=True, color='0F172A')
    ws.merge_cells('B13:E13')
    id_c = ws.cell(row=13, column=2, value=str(customer_dict.get('code', '')))
    id_c.font = Font(name='Arial', size=10, bold=True, color='0F172A')
    id_c.alignment = Alignment(horizontal='left', vertical='center')

    ws.cell(row=13, column=7, value='Fecha:').font = Font(name='Arial', size=10, bold=True, color='0F172A')
    ws.merge_cells('H13:J13')
    f_c = ws.cell(row=13, column=8, value=date_long)
    f_c.font = Font(name='Arial', size=10, color='334155')
    f_c.alignment = Alignment(horizontal='center', vertical='center')

    # Dotted rows for Nombre, Dirección, Teléfono
    def set_client_detail(row, label, val):
        ws.cell(row=row, column=1, value=label).font = Font(name='Arial', size=10, bold=True, color='0F172A')
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=10)
        vc = ws.cell(row=row, column=2, value=val)
        vc.font = Font(name='Arial', size=10, color='1E293B')
        vc.alignment = Alignment(horizontal='left', vertical='center')
        for c in range(2, 11):
            ws.cell(row=row, column=c).border = Border(
                bottom=dotted_border_side,
                right=thin_border_side if c == 10 else None
            )

    set_client_detail(14, 'Nombre:', customer_dict.get('name', ''))
    set_client_detail(15, 'Dirección:', customer_dict.get('address', ''))
    set_client_detail(16, 'Teléfono:', customer_dict.get('phone', ''))

    # --- INVOICE TABLE (Row 19 onwards) ---
    headers = [
        'No.', 'ID Cliente', 'Factura No.', 'Fecha Factura', 'Fecha Vencim.',
        'Total de Factura', 'Status Deuda', '# de Días', 'Abono', 'Saldo Actual'
    ]
    tbl_fill = PatternFill(start_color='CBD5E1', end_color='CBD5E1', fill_type='solid')
    header_font = Font(name='Arial', size=10, bold=True, color='0F172A')

    ws.row_dimensions[19].height = 28
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=19, column=col_idx, value=h)
        cell.fill = tbl_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    curr_row = 20
    data_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    data_font = Font(name='Arial', size=10, color='1E293B')

    for inv in invoices_list:
        ws.row_dimensions[curr_row].height = 20
        # A: No.
        c_a = ws.cell(row=curr_row, column=1, value=inv['no'])
        c_a.alignment = Alignment(horizontal='center', vertical='center')

        # B: ID Cliente
        c_b = ws.cell(row=curr_row, column=2, value=inv['customer_code'])
        c_b.alignment = Alignment(horizontal='center', vertical='center')

        # C: Factura No.
        c_c = ws.cell(row=curr_row, column=3, value=inv['number'])
        c_c.alignment = Alignment(horizontal='center', vertical='center')
        c_c.font = Font(name='Arial', size=10, bold=True, color='0F172A')

        # D: Fecha Factura
        c_d = ws.cell(row=curr_row, column=4, value=inv['issue_date_short'])
        c_d.alignment = Alignment(horizontal='center', vertical='center')

        # E: Fecha Vencim.
        c_e = ws.cell(row=curr_row, column=5, value=inv['due_date_short'])
        c_e.alignment = Alignment(horizontal='center', vertical='center')

        # F: Total de Factura
        c_f = ws.cell(row=curr_row, column=6, value=inv['total'])
        c_f.alignment = Alignment(horizontal='right', vertical='center')
        c_f.number_format = '\"Q \"#,##0.00'

        # G: Status Deuda
        c_g = ws.cell(row=curr_row, column=7, value=inv['status_deuda'])
        c_g.alignment = Alignment(horizontal='center', vertical='center')

        # H: # de Días
        c_h = ws.cell(row=curr_row, column=8, value=inv['days'])
        c_h.alignment = Alignment(horizontal='center', vertical='center')

        # I: Abono
        c_i = ws.cell(row=curr_row, column=9)
        if inv['abono'] > 0:
            c_i.value = inv['abono']
            c_i.number_format = '\"Q \"#,##0.00'
        else:
            c_i.value = 'Q  -'
        c_i.alignment = Alignment(horizontal='right', vertical='center')

        # J: Saldo Actual
        c_j = ws.cell(row=curr_row, column=10, value=inv['balance'])
        c_j.alignment = Alignment(horizontal='right', vertical='center')
        c_j.number_format = '\"Q \"#,##0.00'

        for col_idx in range(1, 11):
            cell = ws.cell(row=curr_row, column=col_idx)
            cell.font = data_font if col_idx != 3 else Font(name='Arial', size=10, bold=True, color='0F172A')
            cell.border = data_border

        curr_row += 1

    # --- CYAN STRIP & FOOTER ---
    curr_row += 1
    ws.row_dimensions[curr_row].height = 14
    ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=10)
    cyan_fill = PatternFill(start_color='00E5FF', end_color='00E5FF', fill_type='solid')
    for col_idx in range(1, 11):
        ws.cell(row=curr_row, column=col_idx).fill = cyan_fill

    curr_row += 1
    ws.row_dimensions[curr_row].height = 20
    ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=3)
    c_h = ws.cell(row=curr_row, column=1, value=f'Hora: {time_str}')
    c_h.font = Font(name='Arial', size=9, bold=True, color='475569')
    c_h.alignment = Alignment(horizontal='left', vertical='center')

    ws.merge_cells(start_row=curr_row, start_column=4, end_row=curr_row, end_column=7)
    c_p = ws.cell(row=curr_row, column=4, value='Página 1 de 1')
    c_p.font = Font(name='Arial', size=9, bold=True, color='475569')
    c_p.alignment = Alignment(horizontal='center', vertical='center')

    ws.merge_cells(start_row=curr_row, start_column=8, end_row=curr_row, end_column=10)
    c_f = ws.cell(row=curr_row, column=8, value=f'Fecha: {date_footer}')
    c_f.font = Font(name='Arial', size=9, bold=True, color='475569')
    c_f.alignment = Alignment(horizontal='right', vertical='center')

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get("/customers/{customer_id}/export-excel")
async def export_customer_statement_excel(
    customer_id: str,
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Descargar estado de cuenta de cliente en formato Excel (.xlsx) con el diseño oficial."""
    c_uuid = uuid.UUID(str(current_user.company_id))
    await _sync_sales_to_credits(db, c_uuid)

    try:
        cust_uuid = uuid.UUID(customer_id)
    except ValueError:
        raise HTTPException(status_code=400, detail="ID de cliente inválido")

    c_stmt = select(Partner).where(
        Partner.id == cust_uuid,
        Partner.company_id == c_uuid
    )
    c_res = await db.execute(c_stmt)
    customer = c_res.scalar_one_or_none()
    if not customer:
        raise HTTPException(status_code=404, detail="Cliente no encontrado")

    comp_stmt = select(Company).where(Company.id == c_uuid)
    comp_res = await db.execute(comp_stmt)
    company = comp_res.scalar_one_or_none()

    inv_stmt = select(CustomerInvoiceCredit).where(
        CustomerInvoiceCredit.customer_id == cust_uuid,
        CustomerInvoiceCredit.company_id == c_uuid,
        CustomerInvoiceCredit.is_deleted == False
    ).order_by(CustomerInvoiceCredit.issue_date.asc())
    inv_res = await db.execute(inv_stmt)
    invoices = inv_res.scalars().all()

    now_dt = datetime.now()
    today_date = now_dt.date()
    dias_semana = ['lunes', 'martes', 'miércoles', 'jueves', 'viernes', 'sábado', 'domingo']
    meses_nombre = ['enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre']
    date_long = f"{dias_semana[now_dt.weekday()]}, {now_dt.day} de {meses_nombre[now_dt.month - 1]} de {now_dt.year}"
    date_footer = f"{now_dt.day} de {meses_nombre[now_dt.month - 1]} de {now_dt.year}"
    time_str = now_dt.strftime("%H:%M")

    total_envio = 0.0
    total_debt = 0.0
    total_vencido = 0.0
    total_abono = 0.0
    invoices_list = []

    for idx, i in enumerate(invoices, start=1):
        inv_tot = float(i.total or 0)
        inv_bal = float(i.balance or 0)
        inv_abono = max(0.0, inv_tot - inv_bal)

        if i.due_date:
            due_d = i.due_date.date() if hasattr(i.due_date, 'date') else i.due_date
            days_diff = (due_d - today_date).days
        else:
            days_diff = 0

        status_deuda = "Vencido" if days_diff < 0 else "No vencido"

        total_envio += inv_tot
        total_debt += inv_bal
        total_abono += inv_abono
        if days_diff < 0 and inv_bal > 0.001:
            total_vencido += inv_bal

        invoices_list.append({
            "no": idx,
            "id": str(i.id),
            "customer_code": str(customer.code) if customer.code else str(customer.id)[:8],
            "number": i.number,
            "issue_date_short": i.issue_date.strftime("%d-%m-%y") if i.issue_date else "",
            "due_date_short": i.due_date.strftime("%d-%m-%y") if i.due_date else "",
            "total": inv_tot,
            "balance": inv_bal,
            "abono": inv_abono,
            "days": days_diff,
            "status_deuda": status_deuda,
        })

    customer_dict = {
        "code": str(customer.code) if customer.code else str(customer.id)[:8],
        "name": customer.name,
        "address": customer.address or "",
        "phone": customer.phone or customer.mobile or "",
    }
    company_dict = {
        "name": company.name if company and company.name else "Productos Automotrices del Sur, S.A.",
        "nit": company.nit if company and company.nit else "11563009-0",
        "address": company.address if company and company.address else "Lote 67, Calle Central, Nueva Concepción, Esc.",
        "phone": company.phone if company and company.phone else "3454-3857",
        "email": company.email if company and company.email else "proausu.sa@gmail.com",
    }
    summary_dict = {
        "total_envio": round(total_envio, 2),
        "total_deuda": round(total_debt, 2),
        "total_vencido": round(total_vencido, 2),
        "total_abono": round(total_abono, 2),
    }

    excel_bytes = _build_statement_excel(
        customer_dict=customer_dict,
        invoices_list=invoices_list,
        company_dict=company_dict,
        summary_dict=summary_dict,
        date_long=date_long,
        date_footer=date_footer,
        time_str=time_str
    )

    clean_name = "".join(c for c in customer.name if c.isalnum() or c in (' ', '_', '-')).strip().replace(' ', '_')
    filename = f"estado_cuenta_{clean_name}_{now_dt.strftime('%Y%m%d')}.xlsx"

    return Response(
        content=excel_bytes,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


# ===== 13. EXPORT CUSTOMER STATEMENT PDF (.PDF) OFICIAL =====
def _build_statement_pdf(customer_dict, invoices_list, company_dict, summary_dict, date_long, date_footer, time_str) -> bytes:
    if not HAS_REPORTLAB:
        raise HTTPException(status_code=500, detail="Módulo reportlab no disponible en el servidor.")

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=letter,
        leftMargin=24,
        rightMargin=24,
        topMargin=24,
        bottomMargin=24
    )
    styles = getSampleStyleSheet()
    story = []

    norm_style = ParagraphStyle(
        'Norm',
        parent=styles['Normal'],
        fontSize=8.5,
        leading=11,
        fontName='Helvetica',
        textColor=rl_colors.HexColor('#0F172A')
    )

    # 1. Top Section (Company card on left, Estado de Cuenta on right)
    banner_path = os.path.join(os.getcwd(), 'frontend', 'static', 'img', 'statement_logos_banner.png')
    left_elements = []
    if os.path.exists(banner_path):
        left_elements.append(RLImage(banner_path, width=220, height=45))
    left_elements.append(Spacer(1, 4))
    left_elements.append(Paragraph(f"<b>{company_dict.get('name', 'Productos Automotrices del Sur, S.A.')}</b>", ParagraphStyle('Comp', parent=norm_style, fontSize=9.5, fontName='Helvetica-Bold')))
    left_elements.append(Paragraph(f"<b>NIT:</b> {company_dict.get('nit', '11563009-0')}", norm_style))
    left_elements.append(Paragraph(f"<b>Dirección:</b> {company_dict.get('address', 'Lote 67, Calle Central, Nueva Concepción, Esc.')}", norm_style))
    left_elements.append(Paragraph(f"<b>Teléfono:</b> {company_dict.get('phone', '3454-3857')}", norm_style))
    left_elements.append(Paragraph(f"<b>Gmail:</b> {company_dict.get('email', 'proausu.sa@gmail.com')}", norm_style))

    left_cell_table = RLTable([[e] for e in left_elements], colWidths=[310])
    left_cell_table.setStyle(RLTableStyle([
        ('BOX', (0, 0), (-1, -1), 1, rl_colors.HexColor('#94A3B8')),
        ('TOPPADDING', (0, 0), (-1, -1), 2),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
    ]))

    tot_envio_fmt = f"Q {summary_dict.get('total_envio', 0.0):,.2f}"
    tot_deuda_fmt = f"Q {summary_dict.get('total_deuda', 0.0):,.2f}"
    tot_vencido_fmt = f"Q {summary_dict.get('total_vencido', 0.0):,.2f}"
    tot_abono_num = summary_dict.get('total_abono', 0.0)
    tot_abono_fmt = f"Q {tot_abono_num:,.2f}" if tot_abono_num > 0 else "Q -"

    right_data = [
        ['Estado de Cuenta', ''],
        ['Total Envío:', tot_envio_fmt],
        ['Total Deuda:', tot_deuda_fmt],
        ['Total Vencido:', tot_vencido_fmt],
        ['Total de Abono:', tot_abono_fmt],
    ]
    right_cell_table = RLTable(right_data, colWidths=[130, 110])
    right_cell_table.setStyle(RLTableStyle([
        ('SPAN', (0, 0), (1, 0)),
        ('BACKGROUND', (0, 0), (1, 0), rl_colors.HexColor('#EEF2F6')),
        ('ALIGN', (0, 0), (1, 0), 'CENTER'),
        ('FONTNAME', (0, 0), (1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (1, 0), 12),
        ('BOTTOMPADDING', (0, 0), (1, 0), 8),
        ('TOPPADDING', (0, 0), (1, 0), 8),
        ('BOX', (0, 0), (-1, -1), 1, rl_colors.HexColor('#94A3B8')),
        ('LINEBELOW', (0, 0), (1, 0), 1, rl_colors.HexColor('#CBD5E1')),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 1), (-1, -1), 9),
        ('ALIGN', (1, 1), (1, -1), 'RIGHT'),
        ('TOPPADDING', (0, 1), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 1), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
    ]))

    top_table = RLTable([[left_cell_table, right_cell_table]], colWidths=[315, 249])
    top_table.setStyle(RLTableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('LEFTPADDING', (0, 0), (-1, -1), 0),
        ('RIGHTPADDING', (0, 0), (-1, -1), 0),
        ('TOPPADDING', (0, 0), (-1, -1), 0),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
    ]))
    story.append(top_table)
    story.append(Spacer(1, 10))

    # 2. Client Card
    client_data = [
        [Paragraph(f"<b>ID Cliente:</b> {customer_dict.get('code', '')}", norm_style), Paragraph(f"<b>Fecha:</b> {date_long}", norm_style)],
        [Paragraph(f"<b>Nombre:</b> {customer_dict.get('name', '')}", norm_style), ''],
        [Paragraph(f"<b>Dirección:</b> {customer_dict.get('address', '') or '-'}", norm_style), ''],
        [Paragraph(f"<b>Teléfono:</b> {customer_dict.get('phone', '') or '-'}", norm_style), ''],
    ]
    client_table = RLTable(client_data, colWidths=[280, 284])
    client_table.setStyle(RLTableStyle([
        ('SPAN', (0, 1), (1, 1)),
        ('SPAN', (0, 2), (1, 2)),
        ('SPAN', (0, 3), (1, 3)),
        ('BOX', (0, 0), (-1, -1), 1, rl_colors.HexColor('#94A3B8')),
        ('ALIGN', (1, 0), (1, 0), 'RIGHT'),
        ('LINEBELOW', (0, 1), (-1, 1), 0.5, rl_colors.HexColor('#CBD5E1')),
        ('LINEBELOW', (0, 2), (-1, 2), 0.5, rl_colors.HexColor('#CBD5E1')),
        ('LINEBELOW', (0, 3), (-1, 3), 0.5, rl_colors.HexColor('#CBD5E1')),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
    ]))
    story.append(client_table)
    story.append(Spacer(1, 10))

    # 3. Invoices Table
    inv_headers = ['No.', 'ID Cliente', 'Factura No.', 'Fecha Factura', 'Fecha Vencim.', 'Total de Factura', 'Status Deuda', '# de Días', 'Abono', 'Saldo Actual']
    inv_rows = [inv_headers]

    for inv in invoices_list:
        tot_f = f"Q {inv['total']:,.2f}"
        abono_f = f"Q {inv['abono']:,.2f}" if inv['abono'] > 0 else "Q -"
        bal_f = f"Q {inv['balance']:,.2f}"
        inv_rows.append([
            str(inv['no']),
            str(inv['customer_code']),
            str(inv['number']),
            str(inv['issue_date_short']),
            str(inv['due_date_short']),
            tot_f,
            str(inv['status_deuda']),
            str(inv['days']),
            abono_f,
            bal_f
        ])

    if len(inv_rows) == 1:
        inv_rows.append(['-', '-', 'Sin facturas a crédito registradas', '-', '-', '-', '-', '-', '-', '-'])

    col_w = [24, 46, 65, 58, 58, 68, 62, 42, 68, 73]
    inv_table = RLTable(inv_rows, colWidths=col_w)
    inv_table.setStyle(RLTableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), rl_colors.HexColor('#CBD5E1')),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 7.5),
        ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
        ('BOX', (0, 0), (-1, -1), 1, rl_colors.HexColor('#94A3B8')),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, rl_colors.HexColor('#E2E8F0')),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('ALIGN', (0, 1), (4, -1), 'CENTER'),
        ('ALIGN', (5, 1), (5, -1), 'RIGHT'),
        ('ALIGN', (6, 1), (7, -1), 'CENTER'),
        ('ALIGN', (8, 1), (9, -1), 'RIGHT'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 2),
        ('RIGHTPADDING', (0, 0), (-1, -1), 3),
    ]))
    story.append(inv_table)
    story.append(Spacer(1, 14))

    # 4. Cyan bar
    cyan_bar = RLTable([['']], colWidths=[564], rowHeights=[10])
    cyan_bar.setStyle(RLTableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), rl_colors.HexColor('#00E5FF')),
        ('TOPPADDING', (0, 0), (-1, -1), 0),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
    ]))
    story.append(cyan_bar)
    story.append(Spacer(1, 4))

    # 5. Bottom Details
    footer_table = RLTable([[f"Hora: {time_str}", "Página 1 de 1", f"Fecha: {date_footer}"]], colWidths=[188, 188, 188])
    footer_table.setStyle(RLTableStyle([
        ('FONTNAME', (0, 0), (-1, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('TEXTCOLOR', (0, 0), (-1, -1), rl_colors.HexColor('#475569')),
        ('ALIGN', (0, 0), (0, 0), 'LEFT'),
        ('ALIGN', (1, 0), (1, 0), 'CENTER'),
        ('ALIGN', (2, 0), (2, 0), 'RIGHT'),
        ('TOPPADDING', (0, 0), (-1, -1), 0),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
    ]))
    story.append(footer_table)

    doc.build(story)
    return buf.getvalue()


@router.get("/customers/{customer_id}/export-pdf")
async def export_customer_statement_pdf(
    customer_id: str,
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Descargar estado de cuenta de cliente en formato PDF (.pdf) con el diseño oficial."""
    c_uuid = uuid.UUID(str(current_user.company_id))
    await _sync_sales_to_credits(db, c_uuid)

    try:
        cust_uuid = uuid.UUID(customer_id)
    except ValueError:
        raise HTTPException(status_code=400, detail="ID de cliente inválido")

    c_stmt = select(Partner).where(
        Partner.id == cust_uuid,
        Partner.company_id == c_uuid
    )
    c_res = await db.execute(c_stmt)
    customer = c_res.scalar_one_or_none()
    if not customer:
        raise HTTPException(status_code=404, detail="Cliente no encontrado")

    comp_stmt = select(Company).where(Company.id == c_uuid)
    comp_res = await db.execute(comp_stmt)
    company = comp_res.scalar_one_or_none()

    inv_stmt = select(CustomerInvoiceCredit).where(
        CustomerInvoiceCredit.customer_id == cust_uuid,
        CustomerInvoiceCredit.company_id == c_uuid,
        CustomerInvoiceCredit.is_deleted == False
    ).order_by(CustomerInvoiceCredit.issue_date.asc())
    inv_res = await db.execute(inv_stmt)
    invoices = inv_res.scalars().all()

    now_dt = datetime.now()
    today_date = now_dt.date()
    dias_semana = ['lunes', 'martes', 'miércoles', 'jueves', 'viernes', 'sábado', 'domingo']
    meses_nombre = ['enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre']
    date_long = f"{dias_semana[now_dt.weekday()]}, {now_dt.day} de {meses_nombre[now_dt.month - 1]} de {now_dt.year}"
    date_footer = f"{now_dt.day} de {meses_nombre[now_dt.month - 1]} de {now_dt.year}"
    time_str = now_dt.strftime("%H:%M")

    total_envio = 0.0
    total_debt = 0.0
    total_vencido = 0.0
    total_abono = 0.0
    invoices_list = []

    for idx, i in enumerate(invoices, start=1):
        inv_tot = float(i.total or 0)
        inv_bal = float(i.balance or 0)
        inv_abono = max(0.0, inv_tot - inv_bal)

        if i.due_date:
            due_d = i.due_date.date() if hasattr(i.due_date, 'date') else i.due_date
            days_diff = (due_d - today_date).days
        else:
            days_diff = 0

        status_deuda = "Vencido" if days_diff < 0 else "No vencido"

        total_envio += inv_tot
        total_debt += inv_bal
        total_abono += inv_abono
        if days_diff < 0 and inv_bal > 0.001:
            total_vencido += inv_bal

        invoices_list.append({
            "no": idx,
            "id": str(i.id),
            "customer_code": str(customer.code) if customer.code else str(customer.id)[:8],
            "number": i.number,
            "issue_date_short": i.issue_date.strftime("%d-%m-%y") if i.issue_date else "",
            "due_date_short": i.due_date.strftime("%d-%m-%y") if i.due_date else "",
            "total": inv_tot,
            "balance": inv_bal,
            "abono": inv_abono,
            "days": days_diff,
            "status_deuda": status_deuda,
        })

    customer_dict = {
        "code": str(customer.code) if customer.code else str(customer.id)[:8],
        "name": customer.name,
        "address": customer.address or "",
        "phone": customer.phone or customer.mobile or "",
    }
    company_dict = {
        "name": company.name if company and company.name else "Productos Automotrices del Sur, S.A.",
        "nit": company.nit if company and company.nit else "11563009-0",
        "address": company.address if company and company.address else "Lote 67, Calle Central, Nueva Concepción, Esc.",
        "phone": company.phone if company and company.phone else "3454-3857",
        "email": company.email if company and company.email else "proausu.sa@gmail.com",
    }
    summary_dict = {
        "total_envio": round(total_envio, 2),
        "total_deuda": round(total_debt, 2),
        "total_vencido": round(total_vencido, 2),
        "total_abono": round(total_abono, 2),
    }

    pdf_bytes = _build_statement_pdf(
        customer_dict=customer_dict,
        invoices_list=invoices_list,
        company_dict=company_dict,
        summary_dict=summary_dict,
        date_long=date_long,
        date_footer=date_footer,
        time_str=time_str
    )

    clean_name = "".join(c for c in customer.name if c.isalnum() or c in (' ', '_', '-')).strip().replace(' ', '_')
    filename = f"estado_cuenta_{clean_name}_{now_dt.strftime('%Y%m%d')}.pdf"

    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
