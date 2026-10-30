"""API v1 — Clientes y Proveedores."""
from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, or_, func, text
from sqlalchemy.orm import selectinload
from pydantic import BaseModel
from typing import Optional, List
from decimal import Decimal
import uuid

from app.core.database import get_db
from app.core.security import get_current_user
from app.models.partner import Partner, PartnerType, CreditStatus
from app.models.sale import SaleOrder, SaleType, SaleState
from app.services.fel_service import lookup_nit_sat

router = APIRouter(prefix="/partners")


def check_partner_management_permission(current_user):
    """Verifica que el usuario no sea solo un cajero sin permisos de gestión de clientes/proveedores."""
    if getattr(current_user, "is_superuser", False) or getattr(current_user, "is_admin", False) or getattr(current_user, "username", "") == "admin" or (getattr(current_user, "email", "") and "admin" in getattr(current_user, "email", "")):
        return
    role_codes = [r.code for r in current_user.roles] if hasattr(current_user, "roles") and current_user.roles else []
    if any(code in role_codes for code in ["admin", "store_manager", "accountant", "inventory_manager", "seller", "preventista", "presale_seller"]):
        return
    if "cashier" in role_codes or not role_codes:
        raise HTTPException(
            status_code=403,
            detail="Acceso denegado: El usuario con rol Cajero no tiene autorización para crear o modificar fichas de clientes."
        )


def check_credit_payment_permission(current_user):
    """Verifica que el usuario tenga autorización para registrar abonos de crédito (Cajeros bloqueados)."""
    if getattr(current_user, "is_superuser", False) or getattr(current_user, "is_admin", False) or getattr(current_user, "username", "") == "admin" or (getattr(current_user, "email", "") and "admin" in getattr(current_user, "email", "")):
        return
    role_codes = [r.code for r in current_user.roles] if hasattr(current_user, "roles") and current_user.roles else []
    if any(code in role_codes for code in ["admin", "store_manager", "accountant"]):
        return
    raise HTTPException(
        status_code=403,
        detail="Acceso denegado: El usuario con rol Cajero no tiene autorización para registrar abonos de crédito."
    )



class PartnerCreateIn(BaseModel):
    code: Optional[int] = None
    name: str
    nit: Optional[str] = "CF"
    cui: Optional[str] = None
    trade_name: Optional[str] = None
    partner_type: str = "customer"
    email: Optional[str] = None
    phone: Optional[str] = None
    mobile: Optional[str] = None
    address: Optional[str] = None
    city: Optional[str] = "Guatemala"
    country: Optional[str] = "Guatemala"
    is_company: bool = False
    credit_status: Optional[str] = "active"
    credit_limit: Decimal = Decimal("0")
    payment_days: int = 0
    image_url: Optional[str] = None
    notes: Optional[str] = None
    tags: Optional[List[str]] = None


class PartnerUpdateIn(BaseModel):
    code: Optional[int] = None
    name: Optional[str] = None
    nit: Optional[str] = None
    cui: Optional[str] = None
    trade_name: Optional[str] = None
    partner_type: Optional[str] = None
    email: Optional[str] = None
    phone: Optional[str] = None
    mobile: Optional[str] = None
    address: Optional[str] = None
    city: Optional[str] = None
    country: Optional[str] = None
    is_company: Optional[bool] = None
    credit_status: Optional[str] = None
    credit_limit: Optional[Decimal] = None
    payment_days: Optional[int] = None
    image_url: Optional[str] = None
    is_active: Optional[bool] = None
    notes: Optional[str] = None
    tags: Optional[List[str]] = None


class CreditPaymentIn(BaseModel):
    amount: Decimal
    payment_method: str = "cash"
    reference: Optional[str] = None
    notes: Optional[str] = None


@router.get("/")
async def list_partners(
    search: Optional[str] = Query(None),
    partner_type: Optional[str] = Query(None),
    type: Optional[str] = Query(None),
    credit_status: Optional[str] = Query(None),
    is_active: Optional[bool] = Query(True),
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Listar y buscar clientes o proveedores."""
    stmt = select(Partner).where(
        Partner.company_id == current_user.company_id,
        Partner.is_deleted == False
    )

    if is_active is not None:
        stmt = stmt.where(Partner.is_active == is_active)

    p_type = partner_type or type
    if p_type:
        try:
            pt_enum = PartnerType(p_type.lower())
            stmt = stmt.where(or_(Partner.partner_type == pt_enum, Partner.partner_type == PartnerType.BOTH))
        except ValueError:
            pass

    if credit_status:
        try:
            cs_enum = CreditStatus(credit_status.lower())
            stmt = stmt.where(Partner.credit_status == cs_enum)
        except ValueError:
            pass

    if search:
        s_clean = search.strip()
        pattern = f"%{s_clean}%"
        conds = [
            Partner.name.ilike(pattern),
            Partner.nit.ilike(pattern),
            Partner.phone.ilike(pattern),
            Partner.email.ilike(pattern),
            Partner.trade_name.ilike(pattern),
        ]
        if s_clean.isdigit():
            conds.append(Partner.code == int(s_clean))
        stmt = stmt.where(or_(*conds))

    stmt = stmt.order_by(Partner.code.asc().nulls_last(), Partner.name.asc())
    res = await db.execute(stmt)
    partners = res.scalars().all()

    items = []
    for p in partners:
        c_limit = Decimal(str(p.credit_limit or 0))
        c_balance = Decimal(str(p.balance or 0))
        c_available = max(Decimal("0"), c_limit - c_balance)
        items.append({
            "id": str(p.id),
            "code": p.code,
            "name": p.name,
            "nit": p.nit or "CF",
            "cui": p.cui or "",
            "trade_name": p.trade_name or "",
            "partner_type": p.partner_type.value if hasattr(p.partner_type, 'value') else str(p.partner_type),
            "email": p.email or "",
            "phone": p.phone or "",
            "mobile": p.mobile or "",
            "address": p.address or "",
            "city": p.city or "Guatemala",
            "credit_limit": float(c_limit),
            "payment_days": p.payment_days or 0,
            "credit_status": p.credit_status.value if hasattr(p.credit_status, 'value') else str(p.credit_status),
            "balance": float(c_balance),
            "available_credit": float(c_available),
            "image_url": p.image_url or "",
            "is_active": p.is_active,
        })

    return {"items": items, "total": len(items)}


@router.get("/lookup-nit/{nit}")
async def lookup_partner_by_nit(
    nit: str,
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """
    Buscar un cliente por NIT para facturación rápida en POS y emisión FEL.
    1. Si es CF: retorna Consumidor Final de inmediato.
    2. Si existe en la base de datos local de la empresa: retorna en 0ms.
    3. Si no existe: consulta en tiempo real ante la SAT / Certificador FEL con validación de Módulo 11.
    """
    nit_raw = nit.strip()
    nit_clean = nit_raw.replace("-", "").replace(" ", "").upper()

    # 1. Consumidor Final
    if nit_clean in ("CF", "CONSUMIDORFINAL", "C/F") or not nit_clean:
        return {
            "found": True,
            "valid": True,
            "nit": "CF",
            "name": "CONSUMIDOR FINAL",
            "trade_name": "Consumidor Final",
            "email": "",
            "address": "CIUDAD DE GUATEMALA",
            "source": "cf",
            "items": [{"nit": "CF", "name": "CONSUMIDOR FINAL", "email": "", "address": "CIUDAD DE GUATEMALA"}]
        }

    # 2. Búsqueda en Base de Datos Local (Prioriza la empresa actual, pero aprovecha catálogo global de clientes)
    stmt = select(Partner).where(
        or_(
            Partner.nit.ilike(nit_raw),
            Partner.nit.ilike(nit_clean),
            func.replace(Partner.nit, "-", "").ilike(nit_clean)
        )
    ).order_by(
        (Partner.company_id == current_user.company_id).desc(),
        Partner.created_at.desc()
    ).limit(1)
    res = await db.execute(stmt)
    partner = res.scalar_one_or_none()

    if partner:
        is_same_company = (partner.company_id == current_user.company_id)
        c_limit = Decimal(str(partner.credit_limit or 0)) if is_same_company else Decimal("0")
        c_balance = Decimal(str(partner.balance or 0)) if is_same_company else Decimal("0")
        c_available = max(Decimal("0"), c_limit - c_balance)
        c_status = partner.credit_status.value if (is_same_company and hasattr(partner.credit_status, 'value')) else (str(partner.credit_status) if is_same_company else "none")
        c_days = partner.payment_days or 0 if is_same_company else 0
        return {
            "found": True,
            "valid": True,
            "id": str(partner.id) if is_same_company else None,
            "nit": partner.nit,
            "name": partner.name,
            "trade_name": partner.trade_name or partner.name,
            "email": partner.email or "",
            "address": partner.address or "CIUDAD DE GUATEMALA",
            "credit_status": c_status,
            "credit_limit": float(c_limit),
            "balance": float(c_balance),
            "available_credit": float(c_available),
            "payment_days": c_days,
            "source": "local_database" if is_same_company else "local_database_shared",
            "items": [{
                "id": str(partner.id) if is_same_company else None,
                "nit": partner.nit,
                "name": partner.name,
                "email": partner.email or "",
                "address": partner.address or "CIUDAD DE GUATEMALA",
                "credit_status": c_status,
                "credit_limit": float(c_limit),
                "balance": float(c_balance),
                "available_credit": float(c_available),
                "payment_days": c_days,
            }]
        }

    # 3. Consulta en tiempo real con algoritmo Módulo 11 SAT y Certificador FEL
    sat_result = await lookup_nit_sat(nit_raw)
    if sat_result.get("valid"):
        formatted_nit = sat_result.get("nit") or nit_raw
        fiscal_name = sat_result.get("name")
        found = bool(fiscal_name)
        address = sat_result.get("address") or "CIUDAD DE GUATEMALA"
        trade_name = sat_result.get("trade_name") or fiscal_name

        return {
            "found": found,
            "valid": True,
            "nit": formatted_nit,
            "name": fiscal_name,
            "trade_name": trade_name,
            "email": "",
            "address": address,
            "source": sat_result.get("source", "sat_api"),
            "message": sat_result.get("message"),
            "items": [{
                "id": None,
                "nit": formatted_nit,
                "name": fiscal_name,
                "email": "",
                "address": address
            }] if fiscal_name else []
        }

    # 4. NIT no superó la validación oficial de Módulo 11 de la SAT
    return {
        "found": False,
        "valid": False,
        "nit": sat_result.get("nit") or nit_raw,
        "name": None,
        "trade_name": None,
        "email": "",
        "address": None,
        "error": sat_result.get("error") or "NIT inválido según algoritmo oficial de SAT",
        "source": "sat_modulo11_failed",
        "items": []
    }


@router.get("/credits")
async def get_customer_credits(
    search: Optional[str] = Query(None),
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Obtener resumen de cuentas por cobrar y créditos de clientes."""
    stmt = select(Partner).where(
        Partner.company_id == current_user.company_id,
        Partner.is_active == True,
        or_(Partner.credit_limit > 0, Partner.balance > 0)
    )
    if search:
        s = f"%{search}%"
        stmt = stmt.where(or_(Partner.name.ilike(s), Partner.nit.ilike(s), Partner.phone.ilike(s)))

    res = await db.execute(stmt)
    partners = res.scalars().all()

    items = []
    total_credit_limit = Decimal("0")
    total_debt = Decimal("0")

    for p in partners:
        c_limit = Decimal(str(p.credit_limit or 0))
        c_balance = Decimal(str(p.balance or 0))
        c_available = max(Decimal("0"), c_limit - c_balance)

        total_credit_limit += c_limit
        total_debt += c_balance

        items.append({
            "id": str(p.id),
            "name": p.name,
            "nit": p.nit,
            "phone": p.phone,
            "email": p.email,
            "credit_limit": float(c_limit),
            "balance": float(c_balance),
            "available_credit": float(c_available),
            "payment_days": p.payment_days or 30,
            "credit_status": p.credit_status.value if hasattr(p.credit_status, "value") else str(p.credit_status)
        })

    return {
        "items": items,
        "total_credit_limit": float(total_credit_limit),
        "total_debt": float(total_debt),
        "total_available": float(max(Decimal("0"), total_credit_limit - total_debt))
    }


@router.get("/kpis")
async def get_partners_kpis(
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Obtener métricas y KPIs ejecutivos de la cartera de clientes y proveedores."""
    stmt = select(Partner).where(Partner.company_id == current_user.company_id, Partner.is_active == True)
    res = await db.execute(stmt)
    partners = res.scalars().all()

    total_customers = sum(1 for p in partners if p.partner_type in (PartnerType.CUSTOMER, PartnerType.BOTH))
    total_suppliers = sum(1 for p in partners if p.partner_type in (PartnerType.SUPPLIER, PartnerType.BOTH))
    total_debt = sum(Decimal(str(p.balance or 0)) for p in partners if (p.balance or 0) > 0)
    total_credit_limit = sum(Decimal(str(p.credit_limit or 0)) for p in partners)
    debtor_customers_count = sum(1 for p in partners if (p.balance or 0) > 0)

    return {
        "total_contacts": len(partners),
        "total_customers": total_customers,
        "total_suppliers": total_suppliers,
        "total_debt": float(total_debt),
        "total_credit_limit": float(total_credit_limit),
        "total_available_credit": float(max(Decimal("0"), total_credit_limit - total_debt)),
        "debtor_customers_count": debtor_customers_count,
    }


# ===== IMPORTACIÓN / EXPORTACIÓN EXCEL Y CSV DE CLIENTES / PROVEEDORES =====

def normalize_partner_media_url(url: str) -> str:
    """Convierte URLs de Google Drive o Dropbox en enlaces directos accesibles para fotos/logos."""
    import re
    if not url:
        return ""
    u = str(url).strip()
    if not u:
        return ""
    if "drive.google.com" in u:
        m = re.search(r"/d/([a-zA-Z0-9_-]+)", u) or re.search(r"[?&]id=([a-zA-Z0-9_-]+)", u)
        if m:
            fid = m.group(1)
            return f"https://drive.google.com/thumbnail?id={fid}&sz=w1000"
    if "dropbox.com" in u:
        if "?dl=0" in u:
            return u.replace("?dl=0", "?raw=1")
        elif "&dl=0" in u:
            return u.replace("&dl=0", "&raw=1")
    return u


PARTNER_TEMPLATE_DEFINITIONS = {
    "codigo": {
        "header": "ID / CODIGO",
        "sample": "1",
        "sample2": "2",
        "description": "Número correlativo/ID del cliente (Opcional, se genera automáticamente si está vacío)."
    },
    "nombre": {
        "header": "NOMBRE / RAZON SOCIAL",
        "sample": "Distribuidora Los Altos, S.A.",
        "sample2": "Juan Carlos Pérez Gómez",
        "description": "Nombre completo o razón social del cliente/proveedor (Obligatorio)."
    },
    "nit": {
        "header": "NIT",
        "sample": "1234567-8",
        "sample2": "CF",
        "description": "NIT guatemalteco o CF (Consumidor Final)."
    },
    "cui": {
        "header": "DPI / CUI",
        "sample": "2500123450101",
        "sample2": "1980123450101",
        "description": "Código Único de Identificación (DPI de 13 dígitos)."
    },
    "tipo": {
        "header": "TIPO DE CONTACTO",
        "sample": "Cliente",
        "sample2": "Proveedor",
        "description": "Cliente, Proveedor, o Ambos."
    },
    "razon_comercial": {
        "header": "NOMBRE COMERCIAL",
        "sample": "Los Altos Auto Partes",
        "sample2": "Taller El Amigo",
        "description": "Nombre de fantasía o comercial."
    },
    "es_empresa": {
        "header": "ES EMPRESA (SI/NO)",
        "sample": "SI",
        "sample2": "NO",
        "description": "SI si es personería jurídica (S.A., Ltda.), NO si es persona individual."
    },
    "telefono": {
        "header": "TELEFONO FIJO",
        "sample": "23456789",
        "sample2": "77661122",
        "description": "Número telefónico PBX o fijo."
    },
    "celular": {
        "header": "CELULAR / WHATSAPP",
        "sample": "55551234",
        "sample2": "44332211",
        "description": "Teléfono móvil o de contacto comercial."
    },
    "email": {
        "header": "CORREO ELECTRONICO",
        "sample": "ventas@losaltos.com",
        "sample2": "juan.perez@gmail.com",
        "description": "Correo para envío de facturas FEL y cotizaciones."
    },
    "direccion": {
        "header": "DIRECCION COMPLETA",
        "sample": "Zona 10, Ciudad de Guatemala",
        "sample2": "3ra Avenida 4-50 Zona 1",
        "description": "Dirección física o fiscal."
    },
    "ciudad": {
        "header": "CIUDAD / MUNICIPIO",
        "sample": "Guatemala",
        "sample2": "Quetzaltenango",
        "description": "Municipio o ciudad de residencia/entrega."
    },
    "departamento_pais": {
        "header": "PAIS",
        "sample": "Guatemala",
        "sample2": "Guatemala",
        "description": "País del contacto."
    },
    "limite_credito": {
        "header": "LIMITE DE CREDITO (Q)",
        "sample": 10000.0,
        "sample2": 0.0,
        "description": "Monto máximo autorizado de crédito en Quetzales."
    },
    "dias_credito": {
        "header": "DIAS DE CREDITO",
        "sample": 30,
        "sample2": 0,
        "description": "Plazo en días para cancelar compras a crédito."
    },
    "estado_credito": {
        "header": "ESTADO DE CREDITO",
        "sample": "Activo",
        "sample2": "Sin Credito",
        "description": "Activo, Bloqueado, Suspendido o Sin Credito."
    },
    "banco": {
        "header": "BANCO",
        "sample": "Banco Industrial",
        "sample2": "BANRURAL",
        "description": "Nombre de la entidad bancaria para transferencias."
    },
    "cuenta_bancaria": {
        "header": "NO. CUENTA BANCARIA",
        "sample": "001-234567-8",
        "sample2": "301-987654-2",
        "description": "Número de cuenta monetaria o de ahorros."
    },
    "imagen_url": {
        "header": "FOTO / LOGO URL",
        "sample": "https://drive.google.com/file/d/1A2B3C.../view",
        "sample2": "",
        "description": "Enlace directo o de Google Drive / Dropbox al logo o foto."
    },
    "notas": {
        "header": "NOTAS / OBSERVACIONES",
        "sample": "Cliente mayorista calificado, descuento del 5%",
        "sample2": "Paga puntual cada quincena",
        "description": "Anotaciones comerciales o preferencias del socio."
    },
}


@router.get("/template-excel")
async def download_partner_excel_template():
    """Descargar plantilla oficial en Microsoft Excel (.xlsx) con estilos Odoo (#714B67) y filas de muestra."""
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from fastapi.responses import Response

    default_cols = [
        "nombre", "nit", "cui", "tipo", "razon_comercial",
        "telefono", "celular", "email", "direccion", "ciudad",
        "limite_credito", "dias_credito", "imagen_url", "notas"
    ]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Directorio Contactos"

    headers = [PARTNER_TEMPLATE_DEFINITIONS[k]["header"] for k in default_cols]
    sample1 = [PARTNER_TEMPLATE_DEFINITIONS[k]["sample"] for k in default_cols]
    sample2 = [PARTNER_TEMPLATE_DEFINITIONS[k]["sample2"] for k in default_cols]

    ws.append(headers)
    ws.append(sample1)
    ws.append(sample2)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="714B67", end_color="714B67", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    ws.row_dimensions[1].height = 28

    output = io.BytesIO()
    wb.save(output)
    xlsx_bytes = output.getvalue()

    return Response(
        content=xlsx_bytes,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="plantilla_clientes_proveedores.xlsx"'}
    )


@router.get("/custom-template-excel")
async def download_custom_partner_excel_template(cols: Optional[str] = Query(None)):
    """Generador dinámico de plantillas en Microsoft Excel (.xlsx) con columnas elegidas por el usuario."""
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from fastapi.responses import Response

    if cols:
        requested = [c.strip().lower() for c in cols.split(",") if c.strip().lower() in PARTNER_TEMPLATE_DEFINITIONS]
    else:
        requested = ["nombre", "nit", "tipo", "telefono", "direccion"]

    if not requested:
        requested = ["nombre", "nit", "tipo", "telefono", "direccion"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Plantilla Personalizada"

    headers = [PARTNER_TEMPLATE_DEFINITIONS[k]["header"] for k in requested]
    sample1 = [PARTNER_TEMPLATE_DEFINITIONS[k]["sample"] for k in requested]
    sample2 = [PARTNER_TEMPLATE_DEFINITIONS[k]["sample2"] for k in requested]

    ws.append(headers)
    ws.append(sample1)
    ws.append(sample2)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="714B67", end_color="714B67", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    ws.row_dimensions[1].height = 28

    output = io.BytesIO()
    wb.save(output)
    xlsx_bytes = output.getvalue()

    return Response(
        content=xlsx_bytes,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="plantilla_clientes_personalizada.xlsx"'}
    )


@router.get("/template-csv")
async def download_partner_csv_template():
    """Descargar plantilla CSV básica para compatibilidad."""
    import csv
    import io
    from fastapi.responses import Response

    output = io.StringIO()
    writer = csv.writer(output, delimiter=",")
    writer.writerow([
        "nombre", "nit", "cui", "tipo", "telefono",
        "email", "direccion", "limite_credito", "dias_credito"
    ])
    writer.writerow([
        "Juan Carlos Pérez", "1234567-8", "2500123450101", "Cliente", "55551234",
        "juan@ejemplo.com", "Zona 10, Ciudad de Guatemala", "5000.00", "30"
    ])
    writer.writerow([
        "Distribuidora El Éxito S.A.", "8765432-1", "", "Proveedor", "22223333",
        "ventas@elexito.com", "Calzada Roosevelt, Mixco", "0.00", "0"
    ])

    csv_data = "\ufeff" + output.getvalue()
    return Response(
        content=csv_data,
        media_type="text/csv",
        headers={"Content-Disposition": 'attachment; filename="plantilla_clientes_proveedores.csv"'}
    )


@router.post("/import-csv")
async def import_partners_csv(
    file: UploadFile = File(...),
    partner_type: str = Query("customer"),
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Importar clientes o proveedores masivamente desde archivo Excel (.xlsx) o CSV con mapeo dinámico de columnas."""
    check_partner_management_permission(current_user)
    import csv
    import io
    import openpyxl

    content = await file.read()
    filename = (file.filename or "").lower()
    is_excel = filename.endswith(".xlsx") or filename.endswith(".xls")

    rows_data = []

    if is_excel:
        try:
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
            sheet = wb.active
            headers = []
            first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
            if not first_row:
                raise HTTPException(status_code=400, detail="El archivo Excel está vacío o no contiene encabezados.")
            for cell in first_row:
                headers.append(str(cell or "").strip())

            for row_idx, row_cells in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                if not any(row_cells):
                    continue
                row_dict = {}
                for col_idx, h in enumerate(headers):
                    if h and col_idx < len(row_cells):
                        val = row_cells[col_idx]
                        row_dict[h] = str(val if val is not None else "").strip()
                rows_data.append((row_idx, row_dict))
        except HTTPException:
            raise
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Error al leer archivo Excel: {e}")
    else:
        text = content.decode("utf-8-sig", errors="replace")
        csv_reader = csv.DictReader(io.StringIO(text))
        if not csv_reader.fieldnames:
            raise HTTPException(status_code=400, detail="El archivo CSV está vacío o mal formateado.")
        for idx, r in enumerate(csv_reader, start=2):
            rows_data.append((idx, r))

    def clean_col(s):
        return str(s).strip().lower().replace(" ", "_").replace("á", "a").replace("é", "e").replace("í", "i").replace("ó", "o").replace("ú", "u").replace("-", "_").replace("/", "_")

    def parse_bool_val(v, default=False):
        if v is None:
            return default
        s = str(v).strip().lower()
        if not s:
            return default
        return s in ["1", "si", "sí", "true", "t", "s", "empresa", "sociedad", "yes", "y"]

    created_count = 0
    updated_count = 0
    errors = []

    # Asegurar que las columnas en PostgreSQL soporten teléfonos largos y nombres extensos
    try:
        from app.core.config import settings
        if "postgresql" in settings.DATABASE_URL:
            await db.execute(text("ALTER TABLE partners ALTER COLUMN phone TYPE VARCHAR(100)"))
            await db.execute(text("ALTER TABLE partners ALTER COLUMN mobile TYPE VARCHAR(100)"))
            await db.execute(text("ALTER TABLE partners ALTER COLUMN nit TYPE VARCHAR(50)"))
            await db.execute(text("ALTER TABLE partners ALTER COLUMN cui TYPE VARCHAR(50)"))
            await db.execute(text("ALTER TABLE partners ALTER COLUMN bank_account TYPE VARCHAR(150)"))
            await db.execute(text("ALTER TABLE partners ALTER COLUMN bank_name TYPE VARCHAR(255)"))
            await db.commit()
    except Exception:
        await db.rollback()

    # Obtener el ID/código correlativo máximo actual de la empresa
    max_code_res = await db.execute(
        select(func.coalesce(func.max(Partner.code), 0)).where(Partner.company_id == current_user.company_id)
    )
    current_max_code = int(max_code_res.scalar() or 0)

    for idx, raw_row in rows_data:
        row = {clean_col(k): str(v).strip() for k, v in raw_row.items() if k}
        row_code_val = (
            row.get("id___codigo")
            or row.get("id_codigo")
            or row.get("codigo")
            or row.get("id")
            or None
        )
        assigned_row_code = None
        if row_code_val:
            try:
                parsed_c = int(str(row_code_val).strip().replace("#", ""))
                if parsed_c > 0:
                    assigned_row_code = parsed_c
            except Exception:
                pass

        name = (
            row.get("nombre")
            or row.get("nombre___razon_social")
            or row.get("nombre_razon_social")
            or row.get("razon_social")
            or row.get("name")
            or row.get("cliente")
            or row.get("proveedor")
        )
        if not name:
            errors.append(f"Fila {idx}: Nombre es obligatorio.")
            continue
        name = str(name).strip()[:300]

        nit = str(row.get("nit") or "CF").strip()[:50]
        raw_cui = row.get("dpi___cui") or row.get("dpi_cui") or row.get("cui") or row.get("dpi") or None
        cui = str(raw_cui).strip()[:50] if raw_cui else None
        
        # Tipo de contacto
        partner_type_str = (
            row.get("tipo_de_contacto")
            or row.get("tipo_contacto")
            or row.get("tipo")
            or row.get("type")
            or "customer"
        ).lower()

        if "amb" in partner_type_str or "both" in partner_type_str:
            p_type = PartnerType.BOTH
        elif "prov" in partner_type_str or "sup" in partner_type_str:
            p_type = PartnerType.SUPPLIER
        else:
            p_type = PartnerType.CUSTOMER

        raw_trade_name = (
            row.get("nombre_comercial")
            or row.get("razon_comercial")
            or row.get("trade_name")
            or None
        )
        trade_name = str(raw_trade_name).strip()[:300] if raw_trade_name else None

        is_company = parse_bool_val(
            row.get("es_empresa__si_no_")
            or row.get("es_empresa")
            or row.get("empresa")
            or row.get("is_company"),
            default=False
        )

        raw_phone = row.get("telefono_fijo") or row.get("telefono") or row.get("phone") or row.get("tel") or None
        phone = str(raw_phone).strip()[:100] if raw_phone else None

        raw_mobile = row.get("celular___whatsapp") or row.get("celular_whatsapp") or row.get("celular") or row.get("mobile") or row.get("whatsapp") or None
        mobile = str(raw_mobile).strip()[:100] if raw_mobile else None

        raw_email = row.get("correo_electronico") or row.get("correo") or row.get("email") or row.get("mail") or None
        email = str(raw_email).strip()[:200] if raw_email else None

        address = row.get("direccion_completa") or row.get("direccion") or row.get("address") or None
        city = str(row.get("ciudad___municipio") or row.get("ciudad_municipio") or row.get("ciudad") or row.get("city") or "Guatemala").strip()[:100]
        country = str(row.get("pais") or row.get("departamento_pais") or row.get("country") or "Guatemala").strip()[:100]

        try:
            credit_limit = Decimal(str(
                row.get("limite_de_credito__q_")
                or row.get("limite_de_credito")
                or row.get("limite_credito")
                or row.get("credit_limit")
                or "0"
            ).replace(",", "."))
        except Exception:
            credit_limit = Decimal("0")

        try:
            payment_days = int(
                row.get("dias_de_credito")
                or row.get("dias_credito")
                or row.get("payment_days")
                or 0
            )
        except Exception:
            payment_days = 0

        # Estado de crédito
        raw_cred_status = (
            row.get("estado_de_credito")
            or row.get("estado_credito")
            or row.get("credit_status")
            or ""
        ).lower()
        if "bloq" in raw_cred_status or "block" in raw_cred_status:
            cred_status = CreditStatus.BLOCKED
        elif "susp" in raw_cred_status:
            cred_status = CreditStatus.SUSPENDED
        elif "act" in raw_cred_status or credit_limit > 0:
            cred_status = CreditStatus.ACTIVE
        else:
            cred_status = CreditStatus.NONE

        raw_bank_name = row.get("banco") or row.get("bank_name") or None
        bank_name = str(raw_bank_name).strip()[:255] if raw_bank_name else None

        raw_bank_account = (
            row.get("no__cuenta_bancaria")
            or row.get("no_cuenta_bancaria")
            or row.get("cuenta_bancaria")
            or row.get("bank_account")
            or None
        )
        bank_account = str(raw_bank_account).strip()[:150] if raw_bank_account else None

        image_url = normalize_partner_media_url(
            row.get("foto___logo_url")
            or row.get("foto_logo_url")
            or row.get("foto_url")
            or row.get("logo_url")
            or row.get("imagen_url")
            or row.get("image_url")
            or row.get("foto")
            or row.get("logo")
            or ""
        )

        notes = (
            row.get("notas___observaciones")
            or row.get("notas_observaciones")
            or row.get("notas")
            or row.get("observaciones")
            or row.get("notes")
            or ""
        )

        existing = None
        if nit and nit.upper() != "CF":
            stmt = select(Partner).where(
                Partner.company_id == current_user.company_id,
                Partner.nit == nit
            )
            res = await db.execute(stmt)
            existing = res.scalars().first()

        if existing:
            if assigned_row_code:
                existing.code = assigned_row_code
            existing.name = name
            if cui:
                existing.cui = cui
            if trade_name:
                existing.trade_name = trade_name
            if p_type:
                existing.partner_type = p_type
            if is_company is not None:
                existing.is_company = is_company
            if phone:
                existing.phone = phone
            if mobile:
                existing.mobile = mobile
            if email:
                existing.email = email
            if address:
                existing.address = address
            if city:
                existing.city = city
            if country:
                existing.country = country
            if credit_limit > 0:
                existing.credit_limit = credit_limit
            if payment_days > 0:
                existing.payment_days = payment_days
            if cred_status != CreditStatus.NONE:
                existing.credit_status = cred_status
            if bank_name:
                existing.bank_name = bank_name
            if bank_account:
                existing.bank_account = bank_account
            if image_url:
                existing.image_url = image_url
            if notes:
                existing.notes = notes
            updated_count += 1
        else:
            if not assigned_row_code:
                current_max_code += 1
                assigned_row_code = current_max_code
            else:
                if assigned_row_code > current_max_code:
                    current_max_code = assigned_row_code

            partner_obj = Partner(
                company_id=current_user.company_id,
                code=assigned_row_code,
                name=name,
                nit=nit,
                cui=cui,
                partner_type=p_type,
                trade_name=trade_name,
                is_company=is_company,
                phone=phone,
                mobile=mobile,
                email=email,
                address=address,
                city=city,
                country=country,
                credit_limit=credit_limit,
                payment_days=payment_days,
                credit_status=cred_status,
                bank_name=bank_name,
                bank_account=bank_account,
                image_url=image_url or None,
                notes=notes,
                is_active=True
            )
            db.add(partner_obj)
            created_count += 1

    try:
        await db.commit()
    except Exception as e:
        await db.rollback()
        raise HTTPException(status_code=400, detail=f"Error al guardar los contactos en la base de datos: {str(e)}")

    return {
        "status": "success",
        "created": created_count,
        "updated": updated_count,
        "errors": errors,
        "total_processed": created_count + updated_count
    }


@router.get("/export-excel")
async def export_partners_excel(
    cols: Optional[str] = Query(None),
    filter_type: Optional[str] = Query("all"),
    search: Optional[str] = Query(None),
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Exportar listado de clientes y proveedores a Microsoft Excel (.xlsx) con columnas y filtros personalizados."""
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from fastapi.responses import Response

    stmt = select(Partner).where(Partner.company_id == current_user.company_id)

    # Filtros por tipo y estado
    if filter_type == "customers":
        stmt = stmt.where(Partner.partner_type.in_([PartnerType.CUSTOMER, PartnerType.BOTH]))
    elif filter_type == "suppliers":
        stmt = stmt.where(Partner.partner_type.in_([PartnerType.SUPPLIER, PartnerType.BOTH]))
    elif filter_type == "with_debt":
        stmt = stmt.where(Partner.balance > 0)
    elif filter_type == "credit_active":
        stmt = stmt.where(Partner.credit_limit > 0)
    elif filter_type == "companies":
        stmt = stmt.where(Partner.is_company == True)
    elif filter_type == "active":
        stmt = stmt.where(Partner.is_active == True)
    elif filter_type == "archived":
        stmt = stmt.where(Partner.is_active == False)

    # Filtro de búsqueda textual
    if search:
        s = f"%{search.strip().lower()}%"
        stmt = stmt.where(
            or_(
                func.lower(Partner.name).like(s),
                func.lower(Partner.nit).like(s),
                func.lower(Partner.cui).like(s),
                func.lower(Partner.trade_name).like(s),
                func.lower(Partner.phone).like(s),
                func.lower(Partner.email).like(s),
            )
        )

    stmt = stmt.order_by(Partner.name)
    res = await db.execute(stmt)
    partners = res.scalars().all()

    def get_type_label(pt):
        if pt == PartnerType.CUSTOMER:
            return "Cliente"
        elif pt == PartnerType.SUPPLIER:
            return "Proveedor"
        elif pt == PartnerType.BOTH:
            return "Cliente / Proveedor"
        return "Contacto"

    def get_credit_label(cs):
        if cs == CreditStatus.ACTIVE:
            return "Activo"
        elif cs == CreditStatus.BLOCKED:
            return "Bloqueado"
        elif cs == CreditStatus.SUSPENDED:
            return "Suspendido"
        return "Sin Crédito"

    partner_col_defs = {
        "codigo": ("ID", lambda p: p.code or ""),
        "nombre": ("NOMBRE / RAZON SOCIAL", lambda p: p.name or ""),
        "nit": ("NIT", lambda p: p.nit or "CF"),
        "cui": ("DPI / CUI", lambda p: p.cui or ""),
        "tipo": ("TIPO DE CONTACTO", lambda p: get_type_label(p.partner_type)),
        "razon_comercial": ("NOMBRE COMERCIAL", lambda p: p.trade_name or ""),
        "es_empresa": ("ES EMPRESA", lambda p: "SI" if p.is_company else "NO"),
        "telefono": ("TELEFONO FIJO", lambda p: p.phone or ""),
        "celular": ("CELULAR / WHATSAPP", lambda p: p.mobile or ""),
        "email": ("CORREO ELECTRONICO", lambda p: p.email or ""),
        "direccion": ("DIRECCION", lambda p: p.address or ""),
        "ciudad": ("CIUDAD / MUNICIPIO", lambda p: p.city or "Guatemala"),
        "departamento_pais": ("PAIS", lambda p: p.country or "Guatemala"),
        "limite_credito": ("LIMITE DE CREDITO (Q)", lambda p: float(p.credit_limit or 0)),
        "saldo_pendiente": ("SALDO PENDIENTE (Q)", lambda p: float(p.balance or 0)),
        "credito_disponible": ("CREDITO DISPONIBLE (Q)", lambda p: max(0.0, float(p.credit_limit or 0) - float(p.balance or 0))),
        "dias_credito": ("DIAS DE CREDITO", lambda p: int(p.payment_days or 0)),
        "estado_credito": ("ESTADO DE CREDITO", lambda p: get_credit_label(p.credit_status)),
        "banco": ("BANCO", lambda p: p.bank_name or ""),
        "cuenta_bancaria": ("NO. CUENTA BANCARIA", lambda p: p.bank_account or ""),
        "imagen_url": ("FOTO / LOGO URL", lambda p: p.image_url or ""),
        "notas": ("NOTAS / OBSERVACIONES", lambda p: p.notes or ""),
        "estado": ("ESTADO", lambda p: "Activo" if p.is_active else "Inactivo"),
    }

    if cols:
        requested = [c.strip().lower() for c in cols.split(",") if c.strip().lower() in partner_col_defs]
    else:
        requested = ["nombre", "nit", "tipo", "telefono", "direccion", "saldo_pendiente"]

    if not requested:
        requested = ["nombre", "nit", "tipo", "telefono", "direccion", "saldo_pendiente"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Directorio de Contactos"

    headers = [partner_col_defs[k][0] for k in requested]
    ws.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="714B67", end_color="714B67", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    for p in partners:
        row_vals = []
        for k in requested:
            fn = partner_col_defs[k][1]
            try:
                row_vals.append(fn(p))
            except Exception:
                row_vals.append("")
        ws.append(row_vals)

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    ws.row_dimensions[1].height = 28

    output = io.BytesIO()
    wb.save(output)
    xlsx_bytes = output.getvalue()

    return Response(
        content=xlsx_bytes,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="directorio_clientes_proveedores.xlsx"'}
    )


@router.get("/export-csv")
async def export_partners_csv(
    partner_type: Optional[str] = None,
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Exportar listado de clientes y/o proveedores a CSV."""
    import csv
    import io
    from fastapi.responses import Response

    stmt = select(Partner).where(
        Partner.company_id == current_user.company_id,
        Partner.is_active == True
    )
    if partner_type:
        stmt = stmt.where(Partner.partner_type == partner_type)

    stmt = stmt.order_by(Partner.name)
    res = await db.execute(stmt)
    partners = res.scalars().all()

    output = io.StringIO()
    writer = csv.writer(output, delimiter=",")
    writer.writerow([
        "ID", "Nombre", "NIT", "CUI", "Tipo",
        "Teléfono", "Email", "Dirección", "Límite Crédito (Q)", "Saldo Deuda (Q)", "Días Crédito"
    ])

    for p in partners:
        p_type = "Cliente" if p.partner_type == PartnerType.CUSTOMER else "Proveedor"
        writer.writerow([
            str(p.id),
            p.name,
            p.nit or "CF",
            p.cui or "",
            p_type,
            p.phone or "",
            p.email or "",
            p.address or "",
            f"{float(p.credit_limit or 0):.2f}",
            f"{float(p.balance or 0):.2f}",
            p.payment_days or 0
        ])

    csv_data = "\ufeff" + output.getvalue()
    return Response(
        content=csv_data,
        media_type="text/csv",
        headers={"Content-Disposition": 'attachment; filename="socios_comerciales.csv"'}
    )


@router.get("/{partner_id}")
async def get_partner(
    partner_id: str,
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Obtener detalle de un cliente/proveedor."""
    try:
        p_uuid = uuid.UUID(partner_id)
    except ValueError:
        raise HTTPException(status_code=400, detail="ID de socio inválido")

    stmt = select(Partner).where(Partner.id == p_uuid, Partner.company_id == current_user.company_id)
    res = await db.execute(stmt)
    p = res.scalar_one_or_none()

    if not p:
        raise HTTPException(status_code=404, detail="Socio no encontrado")

    c_limit = Decimal(str(p.credit_limit or 0))
    c_balance = Decimal(str(p.balance or 0))

    return {
        "id": str(p.id),
        "code": p.code,
        "name": p.name,
        "nit": p.nit,
        "cui": p.cui,
        "trade_name": p.trade_name,
        "partner_type": p.partner_type.value if hasattr(p.partner_type, 'value') else str(p.partner_type),
        "email": p.email,
        "phone": p.phone,
        "mobile": p.mobile,
        "address": p.address,
        "city": p.city,
        "country": p.country or "Guatemala",
        "is_company": p.is_company,
        "credit_status": p.credit_status.value if hasattr(p.credit_status, 'value') else str(p.credit_status),
        "credit_limit": float(c_limit),
        "balance": float(c_balance),
        "available_credit": float(max(Decimal("0"), c_limit - c_balance)),
        "payment_days": p.payment_days or 0,
        "image_url": p.image_url or "",
        "is_active": p.is_active,
        "notes": p.notes,
        "tags": p.tags or [],
    }


@router.get("/{partner_id}/sales")
async def get_partner_sales_history(
    partner_id: str,
    limit: int = Query(50, ge=1, le=200),
    payment_type: Optional[str] = Query(None),
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Obtener historial de compras del cliente separando ventas en efectivo y a crédito con resumen financiero."""
    try:
        p_uuid = uuid.UUID(partner_id)
    except ValueError:
        raise HTTPException(status_code=400, detail="ID de socio inválido")

    stmt = (
        select(SaleOrder)
        .options(selectinload(SaleOrder.payments))
        .where(
            SaleOrder.customer_id == p_uuid,
            SaleOrder.company_id == current_user.company_id,
            SaleOrder.is_deleted == False
        )
        .order_by(SaleOrder.sale_date.desc(), SaleOrder.created_at.desc())
        .limit(limit)
    )

    res = await db.execute(stmt)
    sales = res.scalars().all()

    items = []
    total_cash = Decimal("0")
    total_credit = Decimal("0")
    count_cash = 0
    count_credit = 0
    total_credit_due = Decimal("0")

    for s in sales:
        total_val = Decimal(str(s.total or 0))
        due_val = Decimal(str(s.amount_due or 0))
        paid_val = Decimal(str(s.amount_paid or 0))

        p_methods = [p.payment_method for p in (s.payments or [])]

        is_presale = bool(
            (s.sale_type and (s.sale_type == SaleType.PRESALE or getattr(s.sale_type, "value", str(s.sale_type)) == "presale"))
            or (s.number and s.number.startswith("PRE-"))
        )

        if is_presale and due_val <= Decimal("0") and paid_val <= Decimal("0"):
            due_val = total_val

        # Una venta es crédito si: tiene saldo pendiente > 0, o es pre-venta, o su tipo es credit, o se pagó vía crédito
        is_credit = bool(
            due_val > Decimal("0")
            or is_presale
            or (s.sale_type and (s.sale_type == SaleType.CREDIT or getattr(s.sale_type, "value", str(s.sale_type)) == "credit"))
            or "credit" in p_methods
        )

        if is_credit:
            p_type = "credit"
            p_label = "Crédito"
            total_credit += total_val
            count_credit += 1
            total_credit_due += due_val
        else:
            p_type = "cash"
            p_label = "Efectivo"
            total_cash += total_val
            count_cash += 1

        if payment_type and payment_type != "all" and p_type != payment_type:
            continue

        items.append({
            "id": str(s.id),
            "number": s.number,
            "sale_date": s.sale_date.strftime("%d/%m/%Y %H:%M") if s.sale_date else s.created_at.strftime("%d/%m/%Y %H:%M"),
            "total": float(total_val),
            "amount_paid": float(paid_val),
            "amount_due": float(due_val),
            "is_credit": is_credit,
            "payment_type": p_type,
            "payment_type_label": p_label,
            "payment_methods": p_methods or (["credit"] if is_credit else ["cash"]),
            "state": s.state.value if hasattr(s.state, "value") else str(s.state),
            "fel_uuid": s.fel_uuid,
            "fel_number": s.fel_number,
        })

    return {
        "items": items,
        "total": len(items),
        "summary": {
            "total_cash": float(total_cash),
            "count_cash": count_cash,
            "total_credit": float(total_credit),
            "count_credit": count_credit,
            "total_credit_due": float(total_credit_due),
        }
    }


@router.get("/{partner_id}/statement")
async def get_partner_statement(
    partner_id: str,
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Generar estado de cuenta consolidado del cliente."""
    try:
        p_uuid = uuid.UUID(partner_id)
    except ValueError:
        raise HTTPException(status_code=400, detail="ID de socio inválido")

    stmt = select(Partner).where(Partner.id == p_uuid, Partner.company_id == current_user.company_id)
    res = await db.execute(stmt)
    p = res.scalar_one_or_none()
    if not p:
        raise HTTPException(status_code=404, detail="Socio no encontrado")

    sales_stmt = select(SaleOrder).where(
        SaleOrder.customer_id == p_uuid,
        SaleOrder.company_id == current_user.company_id,
        SaleOrder.is_deleted == False
    ).order_by(SaleOrder.created_at.desc()).limit(50)
    sales_res = await db.execute(sales_stmt)
    sales = sales_res.scalars().all()

    c_limit = Decimal(str(p.credit_limit or 0))

    # Calcular saldo pendiente real de ventas no canceladas
    active_due = sum(
        Decimal(str(s.amount_due if (s.amount_due is not None and s.amount_due > 0) else (s.total if s.state not in [SaleState.PAID] else 0)))
        for s in sales
        if s.state not in [SaleState.CANCELLED, SaleState.REJECTED]
    )
    c_balance = max(Decimal(str(p.balance or 0)), active_due)

    return {
        "customer": {
            "id": str(p.id),
            "name": p.name,
            "trade_name": p.trade_name,
            "nit": p.nit,
            "cui": p.cui,
            "phone": p.phone or p.mobile,
            "email": p.email,
            "address": p.address,
            "city": p.city,
            "credit_limit": float(c_limit),
            "balance": float(c_balance),
            "available_credit": float(max(Decimal("0"), c_limit - c_balance)),
            "payment_days": p.payment_days or 30,
            "credit_status": p.credit_status.value if hasattr(p.credit_status, "value") else str(p.credit_status),
        },
        "sales": [
            {
                "id": str(s.id),
                "number": s.number,
                "date": s.sale_date.strftime("%d/%m/%Y %H:%M") if s.sale_date else s.created_at.strftime("%d/%m/%Y %H:%M"),
                "total": float(s.total or 0),
                "amount_paid": float(s.amount_paid or 0),
                "amount_due": float(s.amount_due if (s.amount_due is not None and s.amount_due > 0) else (s.total if s.state not in [SaleState.PAID, SaleState.CANCELLED, SaleState.REJECTED] else 0)),
                "state": s.state.value if hasattr(s.state, "value") else str(s.state),
                "fel_uuid": s.fel_uuid,
            }
            for s in sales
        ],
        "summary": {
            "total_purchases": sum(float(s.total or 0) for s in sales if s.state not in [SaleState.CANCELLED, SaleState.REJECTED]),
            "total_pending_due": float(c_balance),
            "credit_utilization_pct": float(min(Decimal("100"), (c_balance / c_limit * 100))) if c_limit > 0 else 0.0,
        }
    }


@router.post("/")
async def create_partner(
    data: PartnerCreateIn,
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Crear un cliente o proveedor."""
    check_partner_management_permission(current_user)
    try:
        pt_enum = PartnerType(data.partner_type.lower())
    except ValueError:
        pt_enum = PartnerType.CUSTOMER

    try:
        cs_enum = CreditStatus(data.credit_status.lower()) if data.credit_status else CreditStatus.ACTIVE
    except ValueError:
        cs_enum = CreditStatus.ACTIVE

    assigned_code = data.code
    if not assigned_code or assigned_code <= 0:
        max_code_res = await db.execute(
            select(func.coalesce(func.max(Partner.code), 0)).where(Partner.company_id == current_user.company_id)
        )
        assigned_code = int(max_code_res.scalar() or 0) + 1

    partner = Partner(
        company_id=current_user.company_id,
        code=assigned_code,
        name=data.name,
        nit=data.nit or "CF",
        cui=data.cui,
        trade_name=data.trade_name,
        partner_type=pt_enum,
        email=data.email,
        phone=data.phone,
        mobile=data.mobile,
        address=data.address,
        city=data.city or "Guatemala",
        country=data.country or "Guatemala",
        is_company=data.is_company,
        credit_status=cs_enum,
        credit_limit=data.credit_limit,
        payment_days=data.payment_days,
        image_url=data.image_url,
        notes=data.notes,
        tags=data.tags or [],
    )
    db.add(partner)
    await db.commit()
    await db.refresh(partner)

    return {"id": str(partner.id), "code": partner.code, "name": partner.name, "nit": partner.nit, "status": "created"}


@router.put("/{partner_id}")
async def update_partner(
    partner_id: str,
    data: PartnerUpdateIn,
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Actualizar datos de cliente/proveedor."""
    check_partner_management_permission(current_user)
    try:
        p_uuid = uuid.UUID(partner_id)
    except ValueError:
        raise HTTPException(status_code=400, detail="ID de socio inválido")

    stmt = select(Partner).where(Partner.id == p_uuid, Partner.company_id == current_user.company_id)
    res = await db.execute(stmt)
    p = res.scalar_one_or_none()

    if not p:
        raise HTTPException(status_code=404, detail="Socio no encontrado")

    if data.code is not None and data.code > 0: p.code = data.code
    if data.name is not None: p.name = data.name
    if data.nit is not None: p.nit = data.nit
    if data.cui is not None: p.cui = data.cui
    if data.trade_name is not None: p.trade_name = data.trade_name
    if data.partner_type is not None:
        try: p.partner_type = PartnerType(data.partner_type.lower())
        except ValueError: pass
    if data.email is not None: p.email = data.email
    if data.phone is not None: p.phone = data.phone
    if data.mobile is not None: p.mobile = data.mobile
    if data.address is not None: p.address = data.address
    if data.city is not None: p.city = data.city
    if data.country is not None: p.country = data.country
    if data.is_company is not None: p.is_company = data.is_company
    if data.credit_status is not None:
        try: p.credit_status = CreditStatus(data.credit_status.lower())
        except ValueError: pass
    if data.credit_limit is not None: p.credit_limit = data.credit_limit
    if data.payment_days is not None: p.payment_days = data.payment_days
    if data.image_url is not None: p.image_url = data.image_url
    if data.is_active is not None: p.is_active = data.is_active
    if data.notes is not None: p.notes = data.notes
    if data.tags is not None: p.tags = data.tags

    await db.commit()
    return {"id": str(p.id), "status": "updated"}


@router.post("/{partner_id}/credit-payments")
async def pay_customer_credit(
    partner_id: str,
    data: CreditPaymentIn,
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Registrar un abono o pago a la deuda del cliente."""
    check_credit_payment_permission(current_user)
    try:
        p_uuid = uuid.UUID(partner_id)
    except ValueError:
        raise HTTPException(status_code=400, detail="ID de socio inválido")

    stmt = select(Partner).where(Partner.id == p_uuid, Partner.company_id == current_user.company_id)
    res = await db.execute(stmt)
    partner = res.scalar_one_or_none()

    if not partner:
        raise HTTPException(status_code=404, detail="Cliente no encontrado")

    if data.amount <= 0:
        raise HTTPException(status_code=400, detail="El monto a abonar debe ser mayor a cero")

    current_balance = Decimal(str(partner.balance or 0))
    new_balance = max(Decimal("0"), current_balance - data.amount)
    partner.balance = new_balance

    await db.commit()

    credit_limit = Decimal(str(partner.credit_limit or 0))
    available = max(Decimal("0"), credit_limit - new_balance)

    return {
        "partner_id": str(partner.id),
        "partner_name": partner.name,
        "amount_paid": float(data.amount),
        "previous_balance": float(current_balance),
        "new_balance": float(new_balance),
        "available_credit": float(available),
        "status": "payment_registered"
    }


@router.delete("/{partner_id}")
async def delete_partner(
    partner_id: str,
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Eliminar cliente o proveedor (soft-delete seguro)."""
    check_partner_management_permission(current_user)
    try:
        p_uuid = uuid.UUID(partner_id)
    except ValueError:
        raise HTTPException(status_code=400, detail="ID de socio inválido")

    stmt = select(Partner).where(Partner.id == p_uuid, Partner.company_id == current_user.company_id)
    res = await db.execute(stmt)
    partner = res.scalar_one_or_none()

    if not partner:
        raise HTTPException(status_code=404, detail="Cliente no encontrado")

    partner.is_deleted = True
    partner.is_active = False
    await db.commit()
    return {
        "success": True,
        "status": "deleted",
        "message": f"Cliente {partner.name} eliminado exitosamente."
    }



